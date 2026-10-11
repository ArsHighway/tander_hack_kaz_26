"""Чтение таблиц из того, что дадут организаторы: папка, архив, xlsx, xls, ods, csv, json, jsonl
и таблица товаров в Word (docx: таблица с колонкой «Наименование»).

xlsx, csv и jsonl читаются потоково: на сотнях тысяч строк память не растёт. Старый Excel (xls,
до 65 тыс. строк) и ods читаются целиком через calamine.
"""

import codecs
import csv
import json
import logging
import re
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from functools import partial
from pathlib import Path
from typing import Any

log = logging.getLogger(__name__)

GRID_SUFFIXES = {".xls", ".xlsb", ".ods"}
TEXT_SUFFIXES = {".csv", ".tsv", ".txt", ".json", ".jsonl", ".ndjson"}
TABLE_SUFFIXES = {".xlsx", ".xlsm", ".docx"} | GRID_SUFFIXES | TEXT_SUFFIXES
ARCHIVE_SUFFIXES = {".zip", ".7z"}
HEADER_SCAN_ROWS = 20
# Число или дата в строке-кандидате значит, что заголовка нет (так выгружены CSV Москвы 2026).
_DATA_CELL = re.compile(r"^[+-]?\d+([.,]\d+)?$|^\d{4}-\d{2}-\d{2}")
# «имя:значение;…», «имя: значение» и длинный текст в заголовке не бывают: такая строка — товар.
_DATA_TEXT = re.compile(r"[^:;]+:[^;]*;|^[^:]{1,60}:\s*\S|^.{61,}$")


@dataclass
class Table:
    source: str
    columns: list[str]
    rows: Callable[[], Iterator[dict[str, Any]]]
    row_count: int | None = None
    # Освобождает файл (книга xlsx в режиме read_only держит его открытым); None — нечего закрывать.
    close: Callable[[], None] | None = None


@dataclass
class SkippedFile:
    path: str
    size: int
    reason: str


def iter_files(root: Path) -> Iterator[Path]:
    """Все файлы под путём; архивы распаковываются рядом в `<имя>.extracted/` (один раз)."""
    if root.is_file():
        if root.suffix.lower() in ARCHIVE_SUFFIXES:
            yield from iter_files(extract_archive(root))
        else:
            yield root
        return
    for path in sorted(root.rglob("*")):
        # Содержимое распакованных архивов придёт через сам архив, второй раз его не читаем.
        inside_archive = any(
            part.endswith((".extracted", ".extracting"))
            for part in path.relative_to(root).parts[:-1]
        )
        if path.is_file() and not inside_archive:
            yield from iter_files(path)


def extract_archive(archive: Path) -> Path:
    target = archive.with_name(archive.name + ".extracted")
    if target.exists():
        return target
    tmp = archive.with_name(archive.name + ".extracting")
    tmp.mkdir(exist_ok=True)
    if archive.suffix.lower() == ".zip":
        with zipfile.ZipFile(archive) as zf:
            zf.extractall(tmp)
    else:
        import py7zr

        with py7zr.SevenZipFile(archive) as sz:
            sz.extractall(tmp)
    tmp.rename(target)
    return target


def read_tables(root: Path) -> tuple[list[Table], list[SkippedFile]]:
    tables: list[Table] = []
    skipped: list[SkippedFile] = []
    for path in iter_files(root):
        suffix = path.suffix.lower()
        if suffix not in TABLE_SUFFIXES:
            skipped.append(
                SkippedFile(
                    str(path), path.stat().st_size, f"не таблица ({suffix or 'без расширения'})"
                )
            )
            continue
        try:
            if suffix in (".xlsx", ".xlsm"):
                tables.extend(_xlsx_tables(path))
            elif suffix in GRID_SUFFIXES:
                tables.extend(_grid_tables(path))
            elif suffix == ".docx":
                tables.append(_docx_table(path))
            elif suffix in (".json",):
                tables.extend(_json_tables(path))
            elif suffix in (".jsonl", ".ndjson"):
                tables.append(_jsonl_table(path))
            else:
                tables.append(_csv_table(path))
        except Exception as exc:  # noqa: BLE001 - файл не прочитан: показываем причину, остальные читаем
            skipped.append(SkippedFile(str(path), path.stat().st_size, _read_error(exc)))
    return tables, skipped


def _read_error(exc: Exception) -> str:
    """Причина для человека: «File is not a zip file» поставщику ничего не говорит."""
    if isinstance(exc, zipfile.BadZipFile):
        return "файл повреждён или это не xlsx: сохраните его заново в Excel или пришлите csv"
    return f"ошибка чтения: {exc}"


def _docx_table(path: Path) -> Table:
    # Импорт здесь: documents сам читает таблицы этим модулем (кольцо импортов на уровне модуля).
    from app.ingest.documents import docx_text, goods_table

    _, tables = docx_text(path.read_bytes())
    rows = goods_table(tables)
    if rows is None:
        raise ValueError("в документе Word нет таблицы товаров с колонкой «Наименование»")
    header = _columns(tuple(rows[0]))
    records = [dict(zip(header, r + [""] * (len(header) - len(r)), strict=False)) for r in rows[1:]]
    return Table(str(path), header, lambda: iter(records), row_count=len(records))


def _xlsx_tables(path: Path) -> Iterator[Table]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        head: list[tuple[Any, ...]] = []
        for row in ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True):
            head.append(row)
        header_idx = next(
            (i for i, row in enumerate(head) if sum(v not in (None, "") for v in row) >= 2), None
        )
        if header_idx is None:
            continue
        has_header = looks_like_header(head[header_idx])
        columns = _columns(head[header_idx], generic=not has_header)
        n = len(columns)
        first_data = header_idx + (2 if has_header else 1)

        def rows(ws=ws, first_data=first_data, columns=columns, n=n) -> Iterator[dict[str, Any]]:
            for row in ws.iter_rows(min_row=first_data, values_only=True):
                cells = row[:n]
                if any(v not in (None, "") for v in cells):
                    yield dict(zip(columns, cells, strict=False))

        count = (ws.max_row - first_data + 1) if ws.max_row else None
        yield Table(f"{path}#{ws.title}", columns, rows, count, close=wb.close)


def _grid_tables(path: Path) -> Iterator[Table]:
    """xls, xlsb, ods: лист целиком в память, заголовок ищется так же, как в xlsx."""
    from python_calamine import CalamineWorkbook

    wb = CalamineWorkbook.from_path(str(path))
    for name in wb.sheet_names:
        grid = [[_cell(v) for v in row] for row in wb.get_sheet_by_name(name).to_python()]
        head = grid[:HEADER_SCAN_ROWS]
        header_idx = next(
            (i for i, row in enumerate(head) if sum(v not in (None, "") for v in row) >= 2), None
        )
        if header_idx is None:
            continue
        has_header = looks_like_header(head[header_idx])
        columns = _columns(tuple(head[header_idx]), generic=not has_header)
        data = [
            dict(zip(columns, row[: len(columns)], strict=False))
            for row in grid[header_idx + (1 if has_header else 0) :]
            if any(v not in (None, "") for v in row[: len(columns)])
        ]
        yield Table(f"{path}#{name}", columns, partial(iter, data), len(data))


def _cell(v: Any) -> Any:
    # В xls все числа — float: «500.0» в характеристике или артикуле поставщик не писал.
    return int(v) if isinstance(v, float) and v.is_integer() else v


def looks_like_header(row: tuple[Any, ...] | list[Any]) -> bool:
    cells = [str(v).strip() for v in row if v not in (None, "")]
    return bool(cells) and not any(_DATA_CELL.match(c) or _DATA_TEXT.search(c) for c in cells)


def _columns(header: tuple[Any, ...], *, generic: bool = False) -> list[str]:
    cells = ["" if v is None else str(v).strip() for v in header]
    while cells and not cells[-1]:
        cells.pop()
    if generic:
        return [f"col_{i + 1}" for i in range(len(cells))]
    seen: dict[str, int] = {}
    columns = []
    for i, name in enumerate(cells):
        name = name or f"col_{i + 1}"
        seen[name] = seen.get(name, 0) + 1
        columns.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return columns


def read_head(path: Path, size: int = 65536) -> bytes:
    with path.open("rb") as fh:
        return fh.read(size)


def detect_encoding(sample: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            # final=False: выборка режет многобайтный символ на границе, это не ошибка кодировки.
            codecs.getincrementaldecoder(enc)().decode(sample, final=False)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def _csv_table(path: Path) -> Table:
    sample = read_head(path)
    encoding = detect_encoding(sample)
    text = sample.decode(encoding, errors="ignore")
    try:
        delimiter = csv.Sniffer().sniff(text, delimiters=";,\t|").delimiter
    except csv.Error:
        delimiter = ";" if text.count(";") > text.count(",") else ","
    with path.open(encoding=encoding, newline="") as fh:
        first = next(csv.reader(fh, delimiter=delimiter), [])
    has_header = looks_like_header(first)
    columns = _columns(tuple(first), generic=not has_header)

    def rows() -> Iterator[dict[str, Any]]:
        with path.open(encoding=encoding, newline="") as fh:
            reader = csv.reader(fh, delimiter=delimiter)
            if has_header:
                next(reader, None)
            for row in reader:
                if any(cell.strip() for cell in row):
                    yield dict(zip(columns, row, strict=False))

    return Table(f"{path} ({encoding}, '{delimiter}')", columns, rows)


def _records_table(source: str, records: list[dict[str, Any]]) -> Table:
    columns: list[str] = []
    for rec in records[:1000]:
        columns.extend(k for k in rec if k not in columns)
    return Table(source, columns, lambda: iter(records), len(records))


def _json_tables(path: Path) -> Iterator[Table]:
    data = json.loads(path.read_text(encoding=detect_encoding(read_head(path))))
    yield from _find_record_lists(str(path), data)


def _find_record_lists(source: str, data: Any, depth: int = 0) -> Iterator[Table]:
    if isinstance(data, list) and data and all(isinstance(x, dict) for x in data[:50]):
        yield _records_table(source, data)
    elif isinstance(data, dict) and depth < 3:
        for key, value in data.items():
            yield from _find_record_lists(f"{source}#{key}", value, depth + 1)


def _jsonl_table(path: Path) -> Table:
    encoding = detect_encoding(read_head(path))
    with path.open(encoding=encoding) as fh:
        first = [json.loads(line) for line, _ in zip(fh, range(1000), strict=False) if line.strip()]
    columns: list[str] = []
    for rec in first:
        columns.extend(k for k in rec if k not in columns)

    def rows() -> Iterator[dict[str, Any]]:
        with path.open(encoding=encoding) as fh:
            for line in fh:
                if line.strip():
                    yield json.loads(line)

    return Table(str(path), columns, rows)
