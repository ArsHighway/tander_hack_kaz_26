"""Массовая загрузка: файл поставщика (xlsx, csv) → те же предложения и проверки по каждой строке.

API разбирает файл, кладёт строки в Redis и ставит в очередь порции по `BATCH_CHUNK_ROWS` строк;
порции разбирают все воркеры (`run_chunk`), результат каждой строки пишется отдельно, порядок в
итоговом файле — исходный. Сбой одной строки не роняет задачу: у строки свой статус и причина.
"""

import asyncio
import contextlib
import io
import logging
import tempfile
import time
from collections.abc import Iterator, Sequence
from datetime import UTC, datetime
from itertools import islice
from pathlib import Path
from typing import Any, Literal, cast
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from pydantic import BaseModel
from redis.asyncio import Redis

from app.checks.detector import detect, verdict
from app.ingest.characteristics import parse_characteristics
from app.ingest.profile import profile_table
from app.ingest.readers import Table, read_tables
from app.jobs import BatchJob
from app.narrowing import question_note
from app.schemas import CardIn, CharIn, Issue, accept_chars

log = logging.getLogger(__name__)

# Предел одного файла (замеры — docs/scaling.md): 100 тыс. строк API разбирает за ~25 с,
# строки занимают ~140 МБ в Redis. Пермская выгрузка (244 тыс.) грузится тремя частями.
MAX_ROWS = 100_000
# 100 тыс. строк пермской выгрузки — ~20 МБ xlsx; csv той же таблицы больше. nginx — с запасом
# (frontend/nginx.conf, client_max_body_size).
MAX_BYTES = 64 * 1024 * 1024
MODULE_COLUMNS = ["Код категории (КПГЗ)", "Код КТРУ", "Заполнено", "Замечания", "Статус"]
STATUS_TEXT = {
    "ready": "готово",
    "needs_answer": "нужен ответ",
    "needs_fix": "нужно исправить",
    "not_found": "нет в справочнике",
    "skipped": "пустая строка",
    "failed": "сбой",
}


class BatchFormatError(ValueError):
    """Файл прочитан, но не похож на список СТЕ."""


class BatchTooLarge(ValueError):
    """Файл больше лимита: такие грузятся частями."""


class BatchRow(BaseModel):
    row: int
    name: str
    description: str = ""
    characteristics: list[CharIn] = []
    # Характеристики, не прошедшие лимиты полей: замечание в результат строки.
    rejected: list[Issue] = []
    empty: bool = False
    raw: dict[str, str] = {}


class BatchRowIn(BaseModel):
    """Что поставщик прислал в строке файла: с этим строка открывается в форме карточки."""

    row: int
    name: str
    description: str
    characteristics: list[CharIn]


RowStatus = Literal["ready", "needs_answer", "needs_fix", "not_found", "skipped", "failed"]


class BatchRowResult(BaseModel):
    row: int
    name: str
    status: RowStatus
    # Итоговый код КТРУ (второй этап по тексту строки); у справочника из позиций КТРУ — код позиции.
    ktru_code: str | None
    # Код категории справочника (КПГЗ) — первый этап; по нему проверены характеристики.
    category_code: str | None = None
    filled: int
    # Из «Заполнено» — подсказки СТЕ портала (такой же или похожих): в тексте значения нет,
    # его проверяет человек.
    hinted: int = 0
    issues: list[Issue]
    # Почему код не подобран (первый шаг сужения) — для «нет в справочнике».
    note: str | None = None
    error: str | None = None
    # Что не сработало у строки: «llm» — характеристики только правилами (сбой или 429 модели),
    # «embeddings» — поиск только по BM25. Под нагрузкой это видно по строкам, а не по логам.
    degraded: list[Literal["llm", "embeddings"]] = []


class BatchState(BaseModel):
    job_id: str
    # cancelled: новые строки не запускаются; начатые до отмены могут сохранить результат.
    status: Literal["queued", "running", "done", "failed", "cancelled"]
    filename: str
    total: int
    processed: int
    # Строк по статусам — счётчики в Redis, без чтения строк: опрос большого пакета дешёвый.
    counts: dict[str, int] = {}
    rows: list[BatchRowResult]
    # Когда файл загружен; у пакетов, созданных до появления поля, — нет.
    created_at: datetime | None = None


class BatchResultPage(BaseModel):
    """Сохранённые строки по всему пакету, в порядке исходного файла.

    Во время обработки новые результаты могут появиться перед курсором: обновление
    начинает просмотр сначала. Необработанные строки в результаты не входят.
    """

    rows: list[BatchRowResult]
    matched: int
    next_after_row: int | None = None


class BatchAccepted(BaseModel):
    job_id: str
    status_url: str


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


@contextlib.contextmanager
def _first_table(filename: str, data: bytes, max_bytes: int | None) -> Iterator[Table]:
    if max_bytes is not None and len(data) > max_bytes:
        raise BatchTooLarge(f"файл больше {max_bytes // (1024 * 1024)} МБ: разбейте его на части")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / Path(filename).name
        path.write_bytes(data)
        tables, skipped = read_tables(Path(tmp))
        if not tables:
            reason = skipped[0].reason if skipped else "таблица не найдена"
            raise BatchFormatError(f"файл не прочитан: {reason}")
        try:
            yield tables[0]
        finally:
            for t in tables:
                if t.close is not None:
                    t.close()


def parse_upload(
    filename: str,
    data: bytes,
    *,
    max_rows: int | None = MAX_ROWS,
    max_bytes: int | None = MAX_BYTES,
) -> list[BatchRow]:
    """Строки файла; лимиты `None` — без лимита (нагрузочный разбор `app.tools.load_test parse`)."""
    with _first_table(filename, data, max_bytes) as table:
        return list(_iter_rows(table, max_rows))


def peek_upload(filename: str, data: bytes, limit: int) -> list[BatchRow]:
    """Первые непустые строки файла, не больше `limit` + 1: «одна карточка или пакет» и «строк
    больше предела» решаются без чтения всего файла."""
    with _first_table(filename, data, MAX_BYTES) as table:
        return list(islice((r for r in _iter_rows(table, None) if not r.empty), limit + 1))


_ROW_JSON = BatchRow.__pydantic_serializer__


def parse_upload_json(
    filename: str,
    data: bytes,
    *,
    max_rows: int | None = MAX_ROWS,
    max_bytes: int | None = MAX_BYTES,
) -> list[bytes]:
    """То же, что `parse_upload`, но каждая строка сразу в JSON: у API в памяти ~1,4 КБ на строку
    вместо ~9 КБ объектов (замер на 100 тыс. строк пермской выгрузки — docs/scaling.md)."""
    with _first_table(filename, data, max_bytes) as table:
        return [_ROW_JSON.to_json(row) for row in _iter_rows(table, max_rows)]


def _iter_rows(table: Table, max_rows: int | None = MAX_ROWS) -> Iterator[BatchRow]:
    """Строки таблицы: роли колонок — по содержимому (`profile_table`), а не по заголовку."""
    profile = profile_table(table)
    if not table.columns or profile.sampled == 0:
        raise BatchFormatError("файл пустой: в нём нет ни одной строки с товаром")
    name_col = profile.column_for("name")
    if name_col is None:
        raise BatchFormatError("не найдена колонка с наименованием СТЕ")
    desc_col = profile.column_for("description")
    char_col = profile.column_for("characteristics")
    for i, record in enumerate(table.rows(), start=1):
        if max_rows is not None and i > max_rows:
            raise BatchTooLarge(f"строк больше {max_rows}: разбейте файл на части")
        name = _text(record.get(name_col))
        chars: list[CharIn] = []
        rejected: list[Issue] = []
        if char_col:
            parsed = parse_characteristics(record.get(char_col))
            chars, rejected = accept_chars([(c.name, c.value.raw, c.unit) for c in parsed.items])
        yield BatchRow(
            row=i,
            name=name,
            description=_text(record.get(desc_col)) if desc_col else "",
            characteristics=chars,
            rejected=rejected,
            empty=not name,
            raw={col: _text(record.get(col)) for col in table.columns},
        )


async def process_row(row: BatchRow, deps: Any) -> BatchRowResult:
    from app.pipeline import AutofillIn, prepare_session, render

    if row.empty:
        return BatchRowResult(
            row=row.row, name=row.name, status="skipped", ktru_code=None, filled=0, issues=[]
        )
    try:
        session = await prepare_session(
            AutofillIn(name=row.name, description=row.description), deps
        )
        out = render(session, deps)
        if not out.ktru.candidates:
            # Позиции нет, но текст поставщика проверяется: мат и данные людей видны и здесь.
            text_issues = [
                *row.rejected,
                *(i for i in out.issues if i.field in ("name", "description")),
            ]
            return BatchRowResult(
                row=row.row, name=row.name, status="not_found", ktru_code=None, filled=0,
                issues=text_issues, note=out.ktru.steps[0] if out.ktru.steps else None,
                degraded=out.degraded,
            )  # fmt: skip
        position = deps.catalog.get(out.ktru.final_code) if out.ktru.final_code else None
        offered = [c for c in out.characteristics if c.value is not None and c.status != "conflict"]
        own = {c.name for c in row.characteristics}
        card = CardIn(
            name=row.name,
            description=row.description,
            characteristics=row.characteristics
            + [
                CharIn(name=c.name, value=c.value or "", unit=c.unit)
                for c in offered
                if c.name not in own
            ],
        )
        # «Такая СТЕ уже есть» — по двойникам, а не по полям карточки: берём из ответа модуля.
        twin = [i for i in out.issues if i.kind == "duplicate"]
        issues = [*row.rejected, *twin, *detect(card, position)]
        # Первый этап выбирает тип/категорию; «готово» требует итоговой позиции второго этапа.
        ktru = out.chain.ktru if out.chain else None
        ktru_code = ktru.final_code if ktru else None
        pending = out.ktru if out.ktru.final_code is None else ktru
        note = None
        if out.ktru.final_code is None or not ktru_code:
            status: Literal["ready", "needs_answer", "needs_fix"] = "needs_answer"
            if pending is not None and pending.question is not None:
                note = question_note(pending.question, pending.steps)
            elif pending is not None and pending.steps:
                note = "; ".join(pending.steps)
            elif out.chain and out.chain.ktru_missing == "no_links":
                note = "Код КТРУ не определён: у выбранной категории нет позиций КТРУ (связи нет)."
            elif out.chain and out.chain.ktru_missing == "not_loaded":
                note = (
                    "Код КТРУ не определён: позиции выбранной категории не загружены в справочник."
                )
            else:
                note = "Код КТРУ не определён. Откройте строку в форме и уточните товар."
        else:
            status = "ready" if verdict(issues) == "ready" else "needs_fix"
        return BatchRowResult(
            row=row.row,
            name=row.name,
            status=status,
            ktru_code=ktru_code,
            category_code=out.ktru.final_code,
            # «Заполнено» — характеристики; поля «Общих сведений» в счёт не идут.
            filled=sum(1 for c in offered if c.group != "general"),
            hinted=sum(
                1 for c in offered if c.method in ("twin", "similar") and c.group != "general"
            ),
            issues=issues,
            note=note,
            degraded=out.degraded,
        )
    except Exception as exc:
        log.exception("строка %d пакета не обработана", row.row)
        return BatchRowResult(
            row=row.row, name=row.name, status="failed", ktru_code=None, filled=0, issues=[],
            error=f"{type(exc).__name__}: {exc}",
        )  # fmt: skip


def _meta_key(job_id: str) -> str:
    return f"batch:{job_id}:meta"


def _rows_key(job_id: str) -> str:
    return f"batch:{job_id}:rows"


def _results_key(job_id: str) -> str:
    return f"batch:{job_id}:results"


def _counts_key(job_id: str) -> str:
    return f"batch:{job_id}:counts"


def _index_keys(job_id: str) -> list[str]:
    return [f"batch:{job_id}:index:{status}" for status in ("all", *STATUS_TEXT)]


# Результат строки одной атомарной операцией: первый результат строки записывается (HSETNX) и
# только он увеличивает счётчик статуса — повторная доставка порции счётчик дважды не увеличит.
# KEYS: results, counts, meta, rows, индексы all и статусов.
# ARGV: строка, результат, статус, TTL обработки, TTL результата, номер ключа индекса статуса.
SAVE_RESULT_LUA = """
local total = redis.call('HGET', KEYS[3], 'total')
if not total then return -1 end
local cancelled = redis.call('HGET', KEYS[3], 'cancelled')
local remaining = redis.call('PTTL', KEYS[3])
if cancelled and remaining <= 0 then return -1 end
local added = redis.call('HSETNX', KEYS[1], ARGV[1], ARGV[2])
local processed = redis.call('HLEN', KEYS[1])
if added == 0 then return processed end
redis.call('HINCRBY', KEYS[2], ARGV[3], 1)
redis.call('ZADD', KEYS[5], ARGV[1], ARGV[1])
redis.call('ZADD', KEYS[tonumber(ARGV[6])], ARGV[1], ARGV[1])
if cancelled then
    for i = 1, #KEYS do redis.call('PEXPIRE', KEYS[i], remaining) end
    return processed
end
local ttl = ARGV[4]
if processed >= tonumber(total) then ttl = ARGV[5] end
for i = 1, #KEYS do redis.call('EXPIRE', KEYS[i], ttl) end
return processed
"""

# Начало порции продлевает жизнь активного пакета, не создавая meta исчезнувшего пакета.
# KEYS: meta, rows, results, counts, индексы; ARGV: TTL обработки.
START_BATCH_LUA = """
local total = redis.call('HGET', KEYS[1], 'total')
if not total or redis.call('HGET', KEYS[1], 'cancelled')
    or redis.call('HLEN', KEYS[3]) >= tonumber(total) then return 0 end
redis.call('HSETNX', KEYS[1], 'started', '1')
for i = 1, #KEYS do redis.call('EXPIRE', KEYS[i], ARGV[1]) end
return 1
"""

# Отмена атомарна относительно старта порции и записи строки.
# Повтор и готовый пакет — без изменений.
# KEYS: meta, rows, results, counts, индексы; ARGV: срок хранения частичного результата.
CANCEL_BATCH_LUA = """
local total = redis.call('HGET', KEYS[1], 'total')
if not total then return -1 end
if redis.call('HGET', KEYS[1], 'cancelled')
    or redis.call('HLEN', KEYS[3]) >= tonumber(total) then return 0 end
redis.call('HSET', KEYS[1], 'cancelled', '1')
for i = 1, #KEYS do redis.call('EXPIRE', KEYS[i], ARGV[1]) end
return 1
"""


# Строк за одну команду RPUSH при создании пакета: большой файл не уходит в Redis одной командой.
_PUSH_SLICE = 1000

# Индекс пакетов для «Моих заявок»: номер пакета → время загрузки. Сам пакет живёт свой срок;
# номер истёкшего убирается при чтении списка. Хранится не больше `BATCH_INDEX_KEEP` последних.
BATCH_INDEX_KEY = "batch:index"
BATCH_INDEX_KEEP = 100


class BatchStore:
    """Пакет в Redis: `meta` (хэш: имя файла, число строк), `rows` (список строк файла в исходном
    порядке), `results` (хэш: номер строки → результат), `counts` (хэш: статус → строк).

    Прогресс — HLEN результатов, счётчики статусов — HINCRBY, оба в одном Lua-скрипте вместе с
    HSETNX результата: повторная доставка порции ничего не увеличит дважды. «Готово» — когда
    результатов столько же, сколько строк; статус выводится при чтении, его не переключают
    наперегонки несколько воркеров. Ожидание и обработка используют `pending_ttl_s`,
    продлеваемый при начале порции и записи новой строки. После последней строки срок
    хранения результата — `ttl_s`. Повторная доставка готового пакета срок не продлевает.
    """

    def __init__(self, redis: Redis, ttl_s: int, pending_ttl_s: int = 86400) -> None:
        self.redis = redis
        self.ttl_s = ttl_s
        self.pending_ttl_s = pending_ttl_s
        self._save = redis.register_script(SAVE_RESULT_LUA)
        self._start = redis.register_script(START_BATCH_LUA)
        self._cancel = redis.register_script(CANCEL_BATCH_LUA)

    def _keys(self, job_id: str) -> tuple[str, str, str]:
        return _meta_key(job_id), _rows_key(job_id), _results_key(job_id)

    async def create(self, filename: str, rows: Sequence[BatchRow | bytes]) -> str:
        """Строки — объекты или готовый JSON (`parse_upload_json`). Строки уходят в Redis кусками
        по `_PUSH_SLICE` (без одной огромной команды), `meta` — последней: пакет «существует» для
        API и воркеров, только когда все строки на месте."""
        job_id = uuid4().hex
        meta, rows_key, _ = self._keys(job_id)
        ttl = self.pending_ttl_s if rows else self.ttl_s
        for i in range(0, len(rows), _PUSH_SLICE):
            part = [
                _ROW_JSON.to_json(r) if isinstance(r, BatchRow) else r
                for r in rows[i : i + _PUSH_SLICE]
            ]
            async with self.redis.pipeline(transaction=False) as pipe:
                pipe.rpush(rows_key, *part)
                pipe.expire(rows_key, ttl)
                await pipe.execute()
        async with self.redis.pipeline(transaction=True) as pipe:
            pipe.hset(
                meta,
                mapping={
                    "job_id": job_id,
                    "filename": filename,
                    "total": len(rows),
                    "result_index": "1",
                    "created_at": datetime.now(UTC).isoformat(),
                },
            )
            pipe.expire(meta, ttl)
            pipe.zadd(BATCH_INDEX_KEY, {job_id: time.time()})
            pipe.zremrangebyrank(BATCH_INDEX_KEY, 0, -BATCH_INDEX_KEEP - 1)
            await pipe.execute()
        return job_id

    async def recent(self, limit: int = 50) -> list[BatchState]:
        """Последние пакеты, новые сверху, только со счётчиками (без строк). Истёкший пакет
        из индекса убирается здесь же: его ключи Redis уже удалил по сроку."""
        ids = cast(list[str], await self.redis.zrevrange(BATCH_INDEX_KEY, 0, limit - 1))
        out: list[BatchState] = []
        gone: list[str] = []
        for job_id in ids:
            state = await self.state(job_id, limit=0)
            if state is None:
                gone.append(job_id)
            else:
                out.append(state)
        if gone:
            await self.redis.zrem(BATCH_INDEX_KEY, *gone)
        return out

    async def total(self, job_id: str) -> int | None:
        raw = await self.redis.hget(_meta_key(job_id), "total")
        return None if raw is None else int(raw)

    async def rows(self, job_id: str, start: int = 0, stop: int | None = None) -> list[BatchRow]:
        """Строки файла с позиции `start` по `stop` (не включая), в исходном порядке."""
        end = -1 if stop is None else stop - 1
        if stop is not None and stop <= start:
            return []
        raw = await self.redis.lrange(_rows_key(job_id), start, end)
        return [BatchRow.model_validate_json(r) for r in raw]

    async def row(self, job_id: str, row: int) -> BatchRow | None:
        """Строка файла по её номеру (с 1, как в результате пакета)."""
        if row < 1:
            return None
        found = await self.rows(job_id, row - 1, row)
        return found[0] if found and found[0].row == row else None

    async def mark_started(self, job_id: str) -> bool:
        meta, rows, results = self._keys(job_id)
        return bool(
            await self._start(
                keys=[meta, rows, results, _counts_key(job_id), *_index_keys(job_id)],
                args=[self.pending_ttl_s],
            )
        )

    async def cancel(self, job_id: str) -> BatchState | None:
        """Запрещает новые строки; результат и входной файл остаются на `ttl_s` секунд.

        Уже начатые строки сохраняют результат, не продлевая срок после отмены.
        Ответ без строк — для дешёвой отмены больших пакетов.
        """
        meta, rows, results = self._keys(job_id)
        found = await self._cancel(
            keys=[meta, rows, results, _counts_key(job_id), *_index_keys(job_id)],
            args=[self.ttl_s],
        )
        return None if int(found) == -1 else await self.state(job_id, limit=0)

    async def can_process_row(self, job_id: str) -> bool:
        """Проверяется после получения слота: ожидающие внутри порции строки тоже отменяются."""
        total, cancelled = await self.redis.hmget(_meta_key(job_id), ["total", "cancelled"])
        return total is not None and cancelled is None

    async def results_for(self, job_id: str, rows: list[int]) -> dict[int, BatchRowResult]:
        if not rows:
            return {}
        found = await self.redis.hmget(_results_key(job_id), [str(r) for r in rows])
        return {
            r: BatchRowResult.model_validate_json(v)
            for r, v in zip(rows, found, strict=True)
            if v is not None
        }

    async def finished_rows(self, job_id: str, rows: list[int]) -> set[int]:
        """Какие из строк уже обработаны (повторная доставка порции их пропускает)."""
        if not rows:
            return set()
        found = await self.redis.hmget(_results_key(job_id), [str(r) for r in rows])
        return {r for r, value in zip(rows, found, strict=True) if value is not None}

    async def save_result(self, job_id: str, result: BatchRowResult) -> int:
        """Атомарно записывает результат; возвращает число результатов, -1 — пакет исчез."""
        meta, rows_key, results = self._keys(job_id)
        keys = [results, _counts_key(job_id), meta, rows_key, *_index_keys(job_id)]
        args: list[str | int] = [
            str(result.row),
            result.model_dump_json(),
            result.status,
            self.pending_ttl_s,
            self.ttl_s,
            6 + list(STATUS_TEXT).index(result.status),
        ]
        return int(await self._save(keys=keys, args=args))

    async def result_page(
        self,
        job_id: str,
        *,
        status: RowStatus | None = None,
        limit: int = 50,
        after_row: int = 0,
    ) -> BatchResultPage | None:
        """Фильтр до лимита. Индекс хранит только номера строк, JSON читается одной страницей.

        Старые пакеты без индекса читаются порциями, без изменения их срока хранения.
        """
        state = await self.state(job_id, limit=0)
        if state is None:
            return None
        matched = state.processed if status is None else state.counts.get(status, 0)
        indexed = await self.redis.hget(_meta_key(job_id), "result_index")
        # Старый воркер мог сохранить часть строк без индекса при обновлении стенда.
        complete_index = await self.redis.zcard(_index_keys(job_id)[0]) >= state.processed
        if indexed and complete_index:
            key = f"batch:{job_id}:index:{status or 'all'}"
            indexed_numbers = cast(
                list[str],
                await self.redis.zrangebyscore(
                    key, f"({after_row}", "+inf", start=0, num=limit + 1
                ),
            )
            found = await self.results_for(job_id, [int(n) for n in indexed_numbers])
            rows = list(found.values())
        else:
            rows = []
            for start in range(after_row + 1, state.total + 1, _PUSH_SLICE):
                numbers = list(range(start, min(start + _PUSH_SLICE, state.total + 1)))
                found = await self.results_for(job_id, numbers)
                rows.extend(r for r in found.values() if status is None or r.status == status)
                if len(rows) > limit:
                    break
        if await self.total(job_id) is None:
            return None
        return BatchResultPage(
            rows=rows[:limit],
            matched=matched,
            next_after_row=rows[limit - 1].row if len(rows) > limit else None,
        )

    async def state(self, job_id: str, limit: int | None = None) -> BatchState | None:
        """Состояние пакета; `limit` — сколько первых строк вернуть (0 — только счётчики)."""
        meta, _, results = self._keys(job_id)
        # decode_responses=True (app.resources): ключи и значения - str.
        info = cast(dict[str, str], await self.redis.hgetall(meta))
        if not info:
            return None
        total = int(info["total"])
        processed = int(await self.redis.hlen(results))
        counts_raw = cast(dict[str, str], await self.redis.hgetall(_counts_key(job_id)))
        counts = {k: int(v) for k, v in counts_raw.items()}
        if limit is None:
            raw = list((await self.redis.hgetall(results)).values())
        elif limit > 0 and total > 0:
            numbers = [str(n) for n in range(1, min(limit, total) + 1)]
            raw = [v for v in await self.redis.hmget(results, numbers) if v is not None]
        else:
            raw = []
        found = sorted((BatchRowResult.model_validate_json(v) for v in raw), key=lambda r: r.row)
        status: Literal["queued", "running", "done", "cancelled"]
        if info.get("cancelled"):
            status = "cancelled"
        elif processed >= total:
            status = "done"
        elif processed or info.get("started"):
            status = "running"
        else:
            status = "queued"
        created = info.get("created_at")
        return BatchState(
            job_id=job_id,
            status=status,
            filename=info.get("filename", ""),
            total=total,
            processed=processed,
            counts=counts,
            rows=found,
            created_at=datetime.fromisoformat(created) if created else None,
        )


async def run_chunk(job: BatchJob, deps: Any, slots: asyncio.Semaphore | None = None) -> int:
    """Порция пакета: строки `job.start:job.stop`; возвращает, сколько строк обработано сейчас.

    Строки порции идут параллельно (`slots` — общий лимит воркера). Уже обработанные строки
    (повторная доставка после падения воркера) пропускаются. Сбой строки — статус `failed` у
    строки; сбой Redis — исключение: брокер вернёт порцию в очередь, другой воркер её доделает.
    """
    store: BatchStore = deps.batches
    if not await store.mark_started(job.job_id):
        log.info("пакет %s завершён, отменён, не найден или истёк", job.job_id)
        return 0
    rows = await store.rows(job.job_id, job.start, job.stop)
    finished = await store.finished_rows(job.job_id, [r.row for r in rows])
    todo = [r for r in rows if r.row not in finished]
    errors: list[Exception] = []
    saved = 0

    async def one(row: BatchRow) -> None:
        nonlocal saved
        try:
            async with slots or contextlib.nullcontext():
                if not await store.can_process_row(job.job_id):
                    return
                result = await process_row(row, deps)
            if await store.save_result(job.job_id, result) >= 0:
                saved += 1
        except Exception as exc:
            # Ошибка записи не отменяет соседние строки. После их завершения брокер
            # вернёт порцию в очередь, а записанные результаты повтор пропустит.
            log.exception("строка %d пакета %s не сохранена", row.row, job.job_id)
            errors.append(exc)

    async with asyncio.TaskGroup() as tg:
        for row in todo:
            tg.create_task(one(row))
    if errors:
        raise ExceptionGroup("не удалось сохранить строки пакета", errors)
    return saved


async def run_batch(job_id: str, deps: Any, slots: asyncio.Semaphore | None = None) -> None:
    """Весь пакет в этом процессе (тесты, отладка); в работе порции раздаёт очередь."""
    await run_chunk(BatchJob(job_id=job_id), deps, slots)


def _text_cell(ws: Any, value: Any) -> Any:
    """Ячейка результата: текст остаётся текстом. openpyxl делает формулой строку с «=», и ячейка
    «=HYPERLINK(…)» из присланного файла стала бы живой формулой в нашем файле."""
    if isinstance(value, str) and value.startswith("="):
        cell = WriteOnlyCell(ws, value=value)
        cell.data_type = "s"
        return cell
    return value


def _columns(rows: list[BatchRow]) -> list[str]:
    return list(rows[0].raw) if rows and rows[0].raw else ["Наименование"]


def _append_rows(
    ws: Any,
    columns: list[str],
    rows: list[BatchRow],
    results: dict[int, BatchRowResult],
    cancelled: bool = False,
) -> None:
    for row in rows:
        res = results.get(row.row)
        base = [row.raw.get(c, "") for c in columns] if row.raw else [row.name]
        if res is None:
            missing = "не обработано — пакет прерван" if cancelled else "не обработано"
            values: list[Any] = [*base, "", "", "", "", missing]
        else:
            notes = "; ".join(
                [
                    *([res.note] if res.note else []),
                    *(f"{i.field}: {i.message}" for i in res.issues if i.severity != "info"),
                    *([f"подсказки СТЕ портала: {res.hinted} — проверьте"] if res.hinted else []),
                ]
            )
            values = [
                *base,
                res.category_code or "",
                res.ktru_code or "",
                res.filled,
                notes,
                STATUS_TEXT[res.status],
            ]
        ws.append([_text_cell(ws, v) for v in values])


def _result_sheet(columns: list[str]) -> tuple[Any, Any]:
    """Потоковая запись (write_only): 100 тыс. строк не держат в памяти дерево ячеек openpyxl."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Результат")
    ws.append(columns + MODULE_COLUMNS)
    return wb, ws


def _saved(wb: Any) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def result_workbook(state: BatchState, rows: list[BatchRow]) -> bytes:
    """xlsx с результатом в исходном порядке строк (всё уже в памяти)."""
    columns = _columns(rows)
    wb, ws = _result_sheet(columns)
    _append_rows(ws, columns, rows, {r.row: r for r in state.rows}, state.status == "cancelled")
    return _saved(wb)


async def result_file(store: BatchStore, job_id: str, *, cancelled: bool = False) -> bytes:
    """xlsx с результатом из Redis кусками по `_PUSH_SLICE` строк: в памяти API один кусок, а не
    весь пакет (100 тыс. строк целиком — ~2 ГБ объектов). Сборка куска — в потоке: цикл
    событий API отвечает другим запросам."""
    total = await store.total(job_id) or 0
    columns = _columns(await store.rows(job_id, 0, 1))
    wb, ws = _result_sheet(columns)
    for start in range(0, total, _PUSH_SLICE):
        rows = await store.rows(job_id, start, start + _PUSH_SLICE)
        results = await store.results_for(job_id, [r.row for r in rows])
        await asyncio.to_thread(_append_rows, ws, columns, rows, results, cancelled)
    return await asyncio.to_thread(_saved, wb)
