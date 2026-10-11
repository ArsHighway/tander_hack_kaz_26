"""Нагрузочный прогон против поднятого стека: `python -m app.tools.load_test <режим> …`.

Режимы:
  subset   первые K строк листа xlsx → отдельный файл (вход остальных режимов);
  check    POST /api/v1/check: RPS и задержки при N одновременных клиентах;
  autofill POST /api/v1/autofill: задержки и доля ответов без LLM (`degraded`);
  batch    POST /api/v1/batch и опрос до `done`: время пакета, строк в секунду, статусы строк;
  parse    разбор файла пакета в этом процессе: время и пик памяти (без лимита строк).

Итог печатается одной JSON-строкой на прогон (и дописывается в `--out`, если задан) — цифры для
`docs/scaling.md`. Скрипт ничего не шлёт, кроме адреса `--base`.
"""

import argparse
import asyncio
import json
import statistics
import sys
import time
from pathlib import Path
from typing import Any

import httpx

from app.batch import BatchRow, parse_upload

# До скольких строк пакета итог читает строки целиком.
FULL_ROWS_MAX = 20000


def _rows(path: Path) -> list[BatchRow]:
    rows = parse_upload(path.name, path.read_bytes(), max_rows=None, max_bytes=None)
    return [r for r in rows if not r.empty]


def _pct(values: list[float], q: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    return round(ordered[min(len(ordered) - 1, int(q * len(ordered)))], 1)


def _report(result: dict[str, Any], out: Path | None) -> None:
    line = json.dumps(result, ensure_ascii=False)
    sys.stdout.write(line + "\n")
    if out is not None:
        with out.open("a", encoding="utf-8") as f:
            f.write(line + "\n")


def subset(src: Path, sheet: str, rows: int, dst: Path) -> None:
    from openpyxl import Workbook, load_workbook

    wb_in = load_workbook(src, read_only=True)
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = sheet
    ws_in = wb_in[sheet] if sheet in wb_in.sheetnames else wb_in.worksheets[0]
    for i, values in enumerate(ws_in.iter_rows(values_only=True)):
        if i > rows:  # заголовок + rows строк
            break
        ws_out.append(list(values))
    wb_out.save(dst)
    wb_in.close()


async def _hammer(
    client: httpx.AsyncClient,
    path: str,
    payloads: list[dict[str, Any]],
    concurrency: int,
    requests: int,
) -> tuple[list[float], list[dict[str, Any]], int, float]:
    """`requests` запросов по кругу из `payloads`, `concurrency` одновременно."""
    latencies: list[float] = []
    bodies: list[dict[str, Any]] = []
    errors = 0
    counter = iter(range(requests))

    async def one_client() -> None:
        nonlocal errors
        for i in counter:
            started = time.perf_counter()
            try:
                resp = await client.post(path, json=payloads[i % len(payloads)])
            except httpx.HTTPError:
                errors += 1
                continue
            latencies.append((time.perf_counter() - started) * 1000)
            if resp.status_code >= 400:
                errors += 1
            else:
                bodies.append(resp.json())

    started = time.perf_counter()
    await asyncio.gather(*(one_client() for _ in range(concurrency)))
    return latencies, bodies, errors, time.perf_counter() - started


async def run_http(
    mode: str, base: str, path: Path, levels: list[int], requests: int, out: Path | None
) -> None:
    rows = _rows(path)
    if mode == "check":
        endpoint = "/api/v1/check"
        payloads = [
            {
                "name": r.name,
                "description": r.description,
                "characteristics": [c.model_dump() for c in r.characteristics],
            }
            for r in rows
        ]
    else:
        endpoint = "/api/v1/autofill"
        payloads = [{"name": r.name, "description": r.description} for r in rows]
    async with httpx.AsyncClient(base_url=base, timeout=300) as client:
        await client.post(endpoint, json=payloads[0])  # прогрев: индекс, соединения
        for level in levels:
            lat, bodies, errors, wall = await _hammer(client, endpoint, payloads, level, requests)
            result: dict[str, Any] = {
                "mode": mode,
                "concurrency": level,
                "requests": requests,
                "errors": errors,
                "wall_s": round(wall, 2),
                "rps": round(len(lat) / wall, 1) if wall else 0.0,
                "p50_ms": _pct(lat, 0.5),
                "p95_ms": _pct(lat, 0.95),
                "p99_ms": _pct(lat, 0.99),
                "mean_ms": round(statistics.fmean(lat), 1) if lat else 0.0,
            }
            if mode == "autofill":
                result["degraded_llm"] = sum("llm" in b.get("degraded", []) for b in bodies)
            _report(result, out)


async def run_batch(base: str, path: Path, timeout_s: float, label: str, out: Path | None) -> None:
    async with httpx.AsyncClient(base_url=base, timeout=600) as client:
        started = time.perf_counter()
        upload = {"file": (path.name, path.read_bytes(), "application/octet-stream")}
        resp = await client.post("/api/v1/batch", files=upload)
        resp.raise_for_status()
        accepted = time.perf_counter() - started
        job_id = resp.json()["job_id"]
        state: dict[str, Any] = {}
        first_progress: float | None = None
        while time.perf_counter() - started < timeout_s:
            state = (await client.get(f"/api/v1/batch/{job_id}?limit=0")).json()
            if first_progress is None and state.get("processed", 0) > 0:
                first_progress = time.perf_counter() - started
            if state.get("status") == "done":
                break
            await asyncio.sleep(0.5)
        wall = time.perf_counter() - started
        # Строки целиком — только у небольших пакетов (ради счёта degraded); у больших — счётчики.
        small = (state.get("total") or 0) <= FULL_ROWS_MAX
        full = (await client.get(f"/api/v1/batch/{job_id}" + ("" if small else "?limit=0"))).json()
    rows = full.get("rows", [])
    statuses = full.get("counts", {})
    work = wall - accepted
    _report(
        {
            "mode": "batch",
            "label": label,
            "job_id": job_id,
            "status": full.get("status"),
            "total": full.get("total"),
            "processed": full.get("processed"),
            "upload_s": round(accepted, 2),
            "first_row_s": round(first_progress, 2) if first_progress is not None else None,
            "wall_s": round(wall, 1),
            "rows_per_s": round((full.get("processed") or 0) / work, 2) if work > 0 else 0.0,
            "statuses": statuses,
            "degraded_llm": sum("llm" in r.get("degraded", []) for r in rows) if small else None,
            "degraded_embeddings": (
                sum("embeddings" in r.get("degraded", []) for r in rows) if small else None
            ),
        },
        out,
    )


def _peak_rss_mb() -> float | None:
    """Пик памяти процесса: Linux (контейнер) — getrusage, Windows — PeakWorkingSetSize."""
    try:
        import resource

        # getrusage нет на Windows (mypy там не видит атрибут); на Windows ImportError ниже
        usage = resource.getrusage(resource.RUSAGE_SELF)  # type: ignore[attr-defined]
        return usage.ru_maxrss / 1024  # Linux: КиБ
    except ImportError:
        pass
    try:
        import ctypes
        from ctypes import wintypes

        class Counters(ctypes.Structure):
            _fields_ = [
                ("cb", wintypes.DWORD),
                ("PageFaultCount", wintypes.DWORD),
                ("PeakWorkingSetSize", ctypes.c_size_t),
                ("WorkingSetSize", ctypes.c_size_t),
                ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
                ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                ("PagefileUsage", ctypes.c_size_t),
                ("PeakPagefileUsage", ctypes.c_size_t),
            ]

        counters = Counters(cb=ctypes.sizeof(Counters))
        windll = ctypes.windll  # type: ignore[attr-defined]
        process = windll.kernel32.GetCurrentProcess()
        if windll.psapi.GetProcessMemoryInfo(process, ctypes.byref(counters), counters.cb):
            return counters.PeakWorkingSetSize / 2**20
    except (AttributeError, OSError):
        pass
    return None


def run_parse(path: Path, out: Path | None) -> None:
    started = time.perf_counter()
    rows = parse_upload(path.name, path.read_bytes(), max_rows=None, max_bytes=None)
    wall = time.perf_counter() - started
    peak_mb = _peak_rss_mb()
    payload = sum(len(r.model_dump_json()) for r in rows)
    _report(
        {
            "mode": "parse",
            "file_mb": round(path.stat().st_size / 2**20, 1),
            "rows": len(rows),
            "wall_s": round(wall, 1),
            "peak_rss_mb": round(peak_mb) if peak_mb is not None else None,
            "rows_json_mb": round(payload / 2**20, 1),
        },
        out,
    )


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    p.add_argument("mode", choices=["subset", "check", "autofill", "batch", "parse"])
    p.add_argument("file", type=Path, help="xlsx/csv со строками СТЕ (для subset — источник)")
    p.add_argument("--base", default="http://localhost:8000")
    p.add_argument("--sheet", default="Все", help="subset: лист источника (нет такого — первый)")
    p.add_argument("--rows", type=int, default=1000, help="subset: сколько строк взять")
    p.add_argument("--dst", type=Path, help="subset: куда записать")
    p.add_argument("--concurrency", default="1,8,32", help="check/autofill: уровни через запятую")
    p.add_argument("--requests", type=int, default=200, help="check/autofill: запросов на уровень")
    p.add_argument("--timeout", type=float, default=3600, help="batch: сколько ждать, с")
    p.add_argument("--label", default="", help="batch: подпись прогона (например, «3 воркера»)")
    p.add_argument("--out", type=Path, help="дописать итог JSON-строкой в файл")
    a = p.parse_args()
    if a.mode == "subset":
        if a.dst is None:
            p.error("subset: нужен --dst")
        subset(a.file, a.sheet, a.rows, a.dst)
    elif a.mode == "parse":
        run_parse(a.file, a.out)
    elif a.mode == "batch":
        asyncio.run(run_batch(a.base, a.file, a.timeout, a.label, a.out))
    else:
        levels = [int(x) for x in a.concurrency.split(",") if x.strip()]
        asyncio.run(run_http(a.mode, a.base, a.file, levels, a.requests, a.out))


if __name__ == "__main__":
    main()
