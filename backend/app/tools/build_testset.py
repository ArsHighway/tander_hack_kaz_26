"""Тестовый набор СТЕ с Портала поставщиков: разные товары, три уровня сложности, xlsx + jsonl.

    uv run python -m app.tools.build_testset --out ../data/testset --total 5000

Источник — тот же публичный поиск каталога, что у сайта (`/newapi/api/Sku/QueryIndex`):
в выдаче уже есть характеристики и описание, отдельный запрос карточки не нужен. Разнообразие
задаёт список запросов по областям (`QUERIES`) и потолки на категорию и на пару «производитель +
категория».

Сложность — насколько характеристики восстановимы из текста поставщика (наименование + описание):
лёгкость = доля значений, найденных в тексте, × min(1, 8 / число характеристик). Набор делится на
три равные части по лёгкости: самая лёгкая треть, средняя, самая трудная.
"""

import argparse
import json
import logging
import random
import re
import time
from collections import Counter
from pathlib import Path
from typing import Any

import httpx
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.portal.cards import PortalCard, parse_card
from app.search.text import tokens

log = logging.getLogger(__name__)

SEARCH_URL = "https://zakupki.mos.ru/newapi/api/Sku/QueryIndex"
# Только ASCII: заголовки HTTP идут в latin-1.
USER_AGENT = "TenderHackKazanTestset/1.0 (+https://zakupki.mos.ru) Mozilla/5.0 (compatible)"
GOODS_PATH = ".1."  # ветка «Товары» дерева КПГЗ
PAGE = 50
PAUSE_S = 0.8
CAP_CATEGORY = 60
CAP_MAKER_CATEGORY = 6
MIN_CHARS = 2

QUERIES: dict[str, list[str]] = {
    "Канцелярия и бумага": [
        "бумага офисная",
        "ручка шариковая",
        "папка",
        "степлер",
        "тетрадь",
        "конверт",
        "маркер",
        "клей канцелярский",
        "скрепки",
        "бумага для заметок",
        "калькулятор",
    ],
    "Оргтехника и IT": [
        "картридж",
        "ноутбук",
        "монитор",
        "клавиатура",
        "мышь компьютерная",
        "принтер",
        "флеш-накопитель",
        "жесткий диск",
        "коммутатор",
        "кабель патч-корд",
        "сервер",
        "источник бесперебойного питания",
        "проектор",
    ],
    "Хозяйственные товары": [
        "мешки для мусора",
        "перчатки",
        "моющее средство",
        "бумага туалетная",
        "полотенца бумажные",
        "швабра",
        "ведро",
        "мыло жидкое",
        "салфетки",
        "освежитель воздуха",
        "губка",
    ],
    "Медицина": [
        "перчатки смотровые",
        "шприц",
        "бинт",
        "маска медицинская",
        "катетер",
        "антисептик",
        "тонометр",
        "пластырь",
        "игла",
        "термометр медицинский",
        "реагент",
    ],
    "Электрика и свет": [
        "лампа светодиодная",
        "светильник",
        "кабель силовой",
        "автоматический выключатель",
        "розетка",
        "удлинитель",
        "батарейка",
        "аккумулятор",
        "провод",
    ],
    "Строительство и ремонт": [
        "краска",
        "труба стальная",
        "цемент",
        "саморез",
        "плитка",
        "доска обрезная",
        "смеситель",
        "кран шаровой",
        "утеплитель",
        "герметик",
        "дверь",
    ],
    "Мебель": ["стул офисный", "стол письменный", "шкаф", "кресло", "стеллаж", "тумба", "диван"],
    "Продукты питания": [
        "молоко",
        "масло сливочное",
        "крупа",
        "сахар",
        "мясо говядина",
        "сок",
        "хлеб",
        "консервы",
        "чай",
        "макаронные изделия",
        "яйцо куриное",
    ],
    "Одежда и СИЗ": [
        "костюм рабочий",
        "обувь",
        "халат",
        "респиратор",
        "каска защитная",
        "перчатки рабочие",
        "куртка",
    ],
    "Бытовая техника": [
        "холодильник",
        "чайник электрический",
        "микроволновая печь",
        "пылесос",
        "кондиционер",
        "обогреватель",
        "стиральная машина",
    ],
    "Инструменты": [
        "дрель",
        "отвертка",
        "молоток",
        "набор инструментов",
        "рулетка измерительная",
        "пила",
        "шуруповерт",
    ],
    "Спорт и учёба": [
        "мяч",
        "тренажер",
        "учебник",
        "глобус",
        "доска магнитно-маркерная",
        "игрушка",
        "конструктор",
    ],
    "Автотовары": [
        "шина",
        "масло моторное",
        "аккумулятор автомобильный",
        "фильтр масляный",
        "щетка стеклоочистителя",
    ],
}

_NUM = re.compile(r"\d+(?:[.,]\d+)?")


def fetch(client: httpx.Client, keyword: str, skip: int) -> list[dict]:
    query = {
        "take": PAGE,
        "skip": skip,
        "withCount": False,
        "filter": {
            "keyword": keyword,
            "isExact": False,
            "productionDirectoryPaths": {"values": [GOODS_PATH]},
        },
    }
    # ensure_ascii: кириллица в URL ломается у портала молча (выдача не по делу при статусе 200).
    for attempt in range(4):
        try:
            r = client.get(SEARCH_URL, params={"queryFilter": json.dumps(query, ensure_ascii=True)})
            if r.status_code == httpx.codes.OK:
                return r.json().get("items") or []
            log.warning("%s skip=%s: HTTP %s", keyword, skip, r.status_code)
        except httpx.HTTPError as e:
            log.warning("%s skip=%s: %s", keyword, skip, e)
        time.sleep(2 * (attempt + 1))
    return []


def to_record(item: dict, root: str) -> dict:
    """Позиция выдачи → строка `data/portal/*.jsonl` (формат `app.portal.cards.parse_card`)."""
    chars = []
    for row in item.get("skuCharacteristics") or []:
        meta = row.get("productCharacteristicValue") or {}
        value = next(
            (
                v
                for v in (
                    row.get("characteristicValueStringValue"),
                    row.get("characteristicValueDecimalValue"),
                    row.get("characteristicValueIntValue"),
                    row.get("characteristicValueBoolValue"),
                    row.get("rangeValue"),
                )
                if v not in (None, "")
            ),
            None,
        )
        if isinstance(value, dict):  # диапазон
            value = (
                "–".join(str(value[k]) for k in ("start", "end") if value.get(k) is not None)
                or None
            )
        chars.append(
            {
                "name": meta.get("productCharacteristicName"),
                "value": value,
                "unit": meta.get("productCharacteristicUnitDesignationNational")
                or meta.get("productCharacteristicNewOkeiName"),
            }
        )
    return {
        "id": item["id"],
        "name": item.get("name"),
        "description": item.get("description") or "",
        "root": root,
        "category": item.get("productionDirectoryName") or "",
        "category_path": item.get("productionDirectoryPath") or "",
        "okpd": item.get("productionCode") or "",
        "manufacturer": item.get("manufacturerName") or "",
        "unit": item.get("okeiName") or "",
        "characteristics": chars,
    }


def _norm(text: str) -> str:
    return text.lower().replace("ё", "е")


def in_text(value: str, text: str) -> bool:
    """Значение восстановимо из текста: число как число, слова — по основам (стемминг поиска)."""
    v = _norm(value)
    if _NUM.fullmatch(v):
        return float(v.replace(",", ".")) in {
            float(n.replace(",", ".")) for n in _NUM.findall(text)
        }
    if v in ("да", "нет"):
        return False
    words = set(tokens(v))
    return bool(words) and words <= set(tokens(text))


def ease(card: PortalCard) -> tuple[float, float]:
    """Доля значений, найденных в тексте, и лёгкость (доля со штрафом за много характеристик)."""
    text = _norm(card.text)
    n = len(card.characteristics)
    share = sum(in_text(c.value, text) for c in card.characteristics) / n
    return share, share * min(1.0, EASY_CHARS / n)


EASY_CHARS = 8
LEVELS = ("лёгкий", "средний", "сложный")
SHEETS = {"лёгкий": "Лёгкие", "средний": "Средние", "сложный": "Сложные"}
HEAD = [
    "Идентификатор СТЕ",
    "Наименование СТЕ",
    "Описание",
    "Категория",
    "Производитель",
    "характеристики СТЕ",
    "Сложность",
    "Характеристик",
    "Доля значений в тексте",
    "Область",
    "ОКПД2",
    "Единица",
    "Путь КПГЗ",
    "Ссылка",
]


def chars_text(card: PortalCard) -> str:
    """Как в выгрузке организаторов: «имя:значение единица;…»."""
    return ";".join(
        f"{c.name}:{c.value}" + (f" {c.unit}" if c.unit else "") for c in card.characteristics
    )


def _cell(v: object) -> object:
    """Управляющие символы из текста поставщиков xlsx не принимает."""
    return ILLEGAL_CHARACTERS_RE.sub(" ", v) if isinstance(v, str) else v


def write_xlsx(
    path: Path, rows: list[tuple[dict, PortalCard, str, float]], stats: list[list]
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, part in [("Все", rows)] + [
        (SHEETS[lv], [r for r in rows if r[2] == lv]) for lv in LEVELS
    ]:
        ws = wb.create_sheet(title)
        ws.append(HEAD)
        for rec, card, level, share in part:
            ws.append(
                [
                    _cell(v)
                    for v in (
                        card.id,
                        card.name,
                        card.description,
                        card.category,
                        rec["manufacturer"],
                        chars_text(card),
                        level,
                        len(card.characteristics),
                        round(share, 2),
                        card.root,
                        card.okpd,
                        rec["unit"],
                        card.category_path,
                        card.url,
                    )
                ]
            )
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        for i, width in enumerate(
            [12, 50, 30, 30, 25, 70, 10, 12, 12, 22, 14, 10, 30, 40], start=1
        ):
            ws.column_dimensions[get_column_letter(i)].width = width
    ws = wb.create_sheet("Как считали")
    for line in stats:
        ws.append(line)
    ws.column_dimensions["A"].width = 60
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=Path("../data/testset"))
    ap.add_argument("--total", type=int, default=5000)
    ap.add_argument("--pages", type=int, default=3, help="страниц по 50 на запрос")
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--no-fetch", action="store_true", help="взять уже скачанный raw.jsonl")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
    args.out.mkdir(parents=True, exist_ok=True)

    raw_path = args.out / "raw.jsonl"
    seen: set[int] = set()
    n_queries = sum(len(v) for v in QUERIES.values())
    if args.no_fetch:
        seen = {
            json.loads(line)["id"] for line in raw_path.read_text(encoding="utf-8").splitlines()
        }
    else:
        fetch_all(raw_path, seen, args.pages, n_queries)
    if not seen:
        raise SystemExit("портал не вернул ни одной позиции — проверьте сеть и адрес поиска")
    build(args, raw_path, seen, n_queries)


def fetch_all(raw_path: Path, seen: set[int], pages: int, n_queries: int) -> None:
    with (
        httpx.Client(headers={"User-Agent": USER_AGENT}, timeout=60) as client,
        raw_path.open("w", encoding="utf-8") as raw,
    ):
        done = 0
        for root, keywords in QUERIES.items():
            for kw in keywords:
                done += 1
                got = 0
                for page in range(pages):
                    items = fetch(client, kw, page * PAGE)
                    time.sleep(PAUSE_S)
                    for item in items:
                        if item.get("id") in seen:
                            continue
                        seen.add(item["id"])
                        raw.write(json.dumps(to_record(item, root), ensure_ascii=False) + "\n")
                        got += 1
                    if len(items) < PAGE:
                        break
                log.info(
                    "[%d/%d] %s / %s: +%d (всего %d)", done, n_queries, root, kw, got, len(seen)
                )


def build(args: argparse.Namespace, raw_path: Path, seen: set[int], n_queries: int) -> None:
    rng = random.Random(args.seed)
    pool: list[tuple[dict, PortalCard, float, float]] = []
    per_cat: Counter[str] = Counter()
    per_maker: Counter[tuple[str, str]] = Counter()
    dropped: Counter[str] = Counter()
    records = [json.loads(line) for line in raw_path.read_text(encoding="utf-8").splitlines()]
    rng.shuffle(records)
    for rec in records:
        card = parse_card(rec)
        if card is None or len(card.characteristics) < MIN_CHARS:
            dropped["меньше 2 характеристик"] += 1
            continue
        if per_cat[card.category] >= CAP_CATEGORY:
            dropped["потолок категории"] += 1
            continue
        if per_maker[(rec["manufacturer"], card.category)] >= CAP_MAKER_CATEGORY:
            dropped["потолок производителя в категории"] += 1
            continue
        per_cat[card.category] += 1
        per_maker[(rec["manufacturer"], card.category)] += 1
        pool.append((rec, card, *ease(card)))
    if len(pool) < args.total:
        log.warning("в пуле %d позиций, меньше запрошенных %d", len(pool), args.total)
    # Терцили лёгкости: самая лёгкая треть, средняя, самая трудная; при равенстве легче та, где
    # характеристик меньше.
    chosen = sorted(pool[: args.total], key=lambda r: (-r[3], len(r[1].characteristics)))
    third = len(chosen) // 3
    bounds = [0, third, 2 * third, len(chosen)]
    rows, tiers = [], []
    for lv, lo, hi in zip(LEVELS, bounds, bounds[1:], strict=False):
        part = chosen[lo:hi]
        rows += [(rec, card, lv, share) for rec, card, share, _ in part]
        n_mean = sum(len(r[1].characteristics) for r in part) / len(part)
        tiers.append(
            [
                lv,
                f"{len(part)} строк; лёгкость {part[-1][3]:.2f}–{part[0][3]:.2f}; "
                f"в тексте в среднем {sum(r[2] for r in part) / len(part):.0%} значений; "
                f"характеристик в среднем {n_mean:.1f}",
            ]
        )

    jsonl = args.out / "testset.jsonl"
    jsonl.write_text(
        "".join(
            json.dumps(
                {**r[0], "difficulty": r[2], "text_share": round(r[3], 3)}, ensure_ascii=False
            )
            + "\n"
            for r in rows
        ),
        encoding="utf-8",
    )
    roots = Counter(r[1].root for r in rows)
    stats: list[list[Any]] = [
        ["Источник", f"zakupki.mos.ru {SEARCH_URL}, {time.strftime('%d.%m.%Y')}, ветка «Товары»"],
        ["Запросов", n_queries],
        ["Получено позиций (уникальных)", len(seen)],
        *[[f"Отброшено: {k}", v] for k, v in dropped.items()],
        ["В пуле после потолков", len(pool)],
        ["В наборе", len(rows)],
        ["Категорий КПГЗ", len({r[1].category for r in rows})],
        [],
        [
            "Сложность",
            "по доле значений характеристик, которые находятся в наименовании и описании",
        ],
        [
            "Лёгкость",
            f"доля значений в тексте × min(1, {EASY_CHARS} / число характеристик); набор делится "
            "на три равные части по лёгкости",
        ],
        *tiers,
        [
            "Значение в тексте",
            "число — то же число; текст — все слова по основам; да/нет — никогда",
        ],
        [],
        ["Область", "строк"],
        *[[k, v] for k, v in roots.most_common()],
    ]
    xlsx = args.out / "testset.xlsx"
    write_xlsx(xlsx, rows, stats)
    log.info("готово: %s, %s; %s", xlsx, jsonl, "; ".join(f"{a}: {b}" for a, b in stats[:8]))


if __name__ == "__main__":
    main()
