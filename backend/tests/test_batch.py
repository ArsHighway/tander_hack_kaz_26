import asyncio
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.v1 import get_deps
from app.batch import (
    MAX_ROWS,
    BatchFormatError,
    BatchRow,
    BatchRowResult,
    BatchStore,
    BatchTooLarge,
    parse_upload,
    parse_upload_json,
    process_row,
    result_file,
    result_workbook,
    run_batch,
    run_chunk,
)
from app.catalog import Catalog
from app.extract import LLMItems
from app.glossary import Glossary
from app.jobs import batch_chunks
from app.main import create_app
from app.pipeline import Deps
from app.schemas import CharIn
from app.search import build_bm25
from app.session import SessionStore
from tests.fakes import FakeLLM, FakeRedis

FAKE = Path(__file__).parent / "fixtures" / "fake"
FORMATS = Path(__file__).parent / "fixtures" / "formats"
CATALOG = Catalog.load([FAKE / "ktru.json"])


def _deps(
    redis: FakeRedis | None = None, llm: FakeLLM | None = None, sent: list | None = None
) -> Deps:
    redis = redis or FakeRedis()

    async def enqueue(job_id: str, total: int) -> None:
        if sent is not None:
            sent.append((job_id, total))

    return Deps(
        catalog=CATALOG,
        bm25=build_bm25(CATALOG),
        vectors=None,
        llm=llm or FakeLLM([LLMItems(items=[])] * 50),  # type: ignore[arg-type]
        glossary=Glossary([]),
        sessions=SessionStore(redis, ttl_s=900),  # type: ignore[arg-type]
        min_score=0.2,
        batches=BatchStore(redis, ttl_s=900),  # type: ignore[arg-type]
        enqueue_batch=enqueue,
    )


def _names(rows: list[BatchRow]) -> list[str]:
    return [r.name for r in rows if not r.empty]


def test_upload_xlsx() -> None:
    rows = parse_upload("ste.xlsx", (FAKE / "ste.xlsx").read_bytes())
    assert len(_names(rows)) > 0


def test_upload_csv_in_cp1251_without_header_gives_the_same_names() -> None:
    xlsx = _names(parse_upload("ste.xlsx", (FAKE / "ste.xlsx").read_bytes()))
    cp1251 = _names(parse_upload("ste_cp1251.csv", (FAKE / "ste_cp1251.csv").read_bytes()))
    headerless = _names(
        parse_upload("ste_headerless.csv", (FAKE / "ste_headerless.csv").read_bytes())
    )
    assert xlsx == cp1251 == headerless


def test_upload_keeps_supplier_characteristics() -> None:
    rows = parse_upload("ste.xlsx", (FAKE / "ste.xlsx").read_bytes())
    assert any(r.characteristics for r in rows)


def test_blank_row_is_kept_as_skipped() -> None:
    # Строка без наименования, но с описанием, не теряется, а помечается пропущенной.
    # Полностью пустые строки читатель таблиц отбрасывает сам: это не данные.
    data = (
        "Наименование;Описание\nБумага офисная А4;\n;без наименования\nПерчатки нитриловые M;\n"
    ).encode()
    rows = parse_upload("cards.csv", data)
    assert [(r.row, r.empty) for r in rows] == [(1, False), (2, True), (3, False)]


def test_file_without_name_column_is_rejected() -> None:
    with pytest.raises(BatchFormatError):
        parse_upload("cards.csv", "Цена;Количество\n100;2\n".encode())


def test_too_many_rows_is_rejected() -> None:
    data = ("Наименование\n" + "Бумага офисная А4\n" * (MAX_ROWS + 1)).encode()
    with pytest.raises(BatchTooLarge, match=str(MAX_ROWS)):
        parse_upload("cards.csv", data)


def test_row_failure_does_not_stop_the_batch() -> None:
    class BrokenBM25:
        def search(self, *_: object, **__: object) -> list:
            raise RuntimeError("индекс сломан")

    deps = _deps()
    deps.bm25 = BrokenBM25()  # type: ignore[assignment]
    result = asyncio.run(process_row(BatchRow(row=1, name="Бумага"), deps))
    assert result.status == "failed"
    assert result.error is not None and "индекс сломан" in result.error


def test_empty_row_is_skipped() -> None:
    result = asyncio.run(process_row(BatchRow(row=2, name="", empty=True), _deps()))
    assert result.status == "skipped"


def test_run_batch_processes_every_row_and_finishes() -> None:
    redis = FakeRedis()
    deps = _deps(redis)
    rows = [
        BatchRow(row=1, name="Батарейка AAA щелочная 1,5 В"),
        BatchRow(row=2, name="", empty=True),
        BatchRow(row=3, name="Услуги по уборке помещений"),
    ]
    store = deps.batches
    assert store is not None
    job_id = asyncio.run(store.create("cards.xlsx", rows))
    asyncio.run(run_batch(job_id, deps))
    state = asyncio.run(store.state(job_id))
    assert state is not None
    assert (state.status, state.total, state.processed) == ("done", 3, 3)
    assert [r.status for r in state.rows][1] == "skipped"
    assert state.rows[2].ktru_code is None


def test_result_workbook_adds_module_columns() -> None:
    rows = [BatchRow(row=1, name="Бумага", raw={"Наименование": "Бумага", "Цена": "100"})]
    from app.batch import BatchRowResult, BatchState

    state = BatchState(
        job_id="x",
        status="done",
        filename="cards.xlsx",
        total=1,
        processed=1,
        rows=[
            BatchRowResult(
                row=1, name="Бумага", status="needs_answer", ktru_code=None, filled=2, issues=[]
            )
        ],
    )
    wb = load_workbook(io.BytesIO(result_workbook(state, rows)))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == [
        "Наименование",
        "Цена",
        "Код категории (КПГЗ)",
        "Код КТРУ",
        "Заполнено",
        "Замечания",
        "Статус",
    ]
    assert ws.max_row == 2


def test_api_accepts_file_and_queues_job() -> None:
    sent: list[str] = []
    redis = FakeRedis()
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps(redis, sent=sent)
    client = TestClient(app)
    data = "Наименование\nБумага офисная А4\n".encode()
    r = client.post("/api/v1/batch", files={"file": ("cards.csv", data, "text/csv")})
    assert r.status_code == 202
    job_id = r.json()["job_id"]
    assert sent == [(job_id, 1)]
    state = client.get(f"/api/v1/batch/{job_id}").json()
    assert (state["status"], state["total"]) == ("queued", 1)
    assert client.get(f"/api/v1/batch/{job_id}/result.xlsx").status_code == 409


def test_api_rejects_oversized_file_with_413() -> None:
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps()
    client = TestClient(app)
    data = ("Наименование\n" + "Бумага офисная А4\n" * (MAX_ROWS + 1)).encode()
    r = client.post("/api/v1/batch", files={"file": ("cards.csv", data, "text/csv")})
    assert r.status_code == 413
    assert str(MAX_ROWS) in r.json()["detail"]


def test_api_reads_no_more_than_the_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    # Загрузка больше предела отвергается, не читаясь в память целиком.
    import app.api.v1 as v1

    monkeypatch.setattr(v1, "MAX_BATCH_BYTES", 100, raising=False)
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps()
    client = TestClient(app)
    data = "\n".join(["Наименование", *["Бумага офисная А4"] * 50]).encode()
    r = client.post("/api/v1/batch", files={"file": ("cards.csv", data, "text/csv")})
    assert r.status_code == 413


def test_api_unknown_batch_is_404() -> None:
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps()
    assert TestClient(app).get("/api/v1/batch/nope").status_code == 404


def test_row_out_of_catalog_is_not_found_with_the_reason() -> None:
    # строка вне справочника была «нужен ответ», хотя спрашивать нечего.
    result = asyncio.run(process_row(BatchRow(row=3, name="Услуги по уборке помещений"), _deps()))
    assert result.status == "not_found"
    assert result.note is not None and result.note.startswith("Категория не найдена")


def test_result_workbook_explains_not_found_rows() -> None:
    from app.batch import BatchRowResult, BatchState

    rows = [BatchRow(row=1, name="Уборка", raw={"Наименование": "Уборка"})]
    note = "Категория не найдена: в справочнике нет подходящей позиции."
    state = BatchState(
        job_id="x",
        status="done",
        filename="cards.xlsx",
        total=1,
        processed=1,
        rows=[
            BatchRowResult(
                row=1, name="Уборка", status="not_found", ktru_code=None, filled=0, issues=[],
                note=note,
            )
        ],
    )  # fmt: skip
    ws = load_workbook(io.BytesIO(result_workbook(state, rows))).active
    assert [c.value for c in ws[2]][-2:] == [note, "нет в справочнике"]


LONG_NAME = "Оценка состояния систем " + "и показателей работы " * 30  # длиннее лимита названия


def _csv(name: str, chars: str) -> bytes:
    return f'Наименование;Характеристики\n{name};"{chars}"\n'.encode()


def test_overlong_characteristic_name_rejects_only_that_characteristic() -> None:
    rows = parse_upload("cards.csv", _csv("Бумага офисная А4", f"Формат:A4;{LONG_NAME}:Да"))
    assert [c.name for c in rows[0].characteristics] == ["Формат"]
    assert [(i.severity, i.field) for i in rows[0].rejected] == [("block", "characteristics")]


def test_overlong_unit_rejects_only_that_characteristic() -> None:
    unit = "количество пикселей в ширину*количество пикселей в высоту"
    record = {
        "Наименование": "Монитор 24",
        "Характеристики": [
            {"name": "Разрешение", "value": "1920x1080", "unit": unit},
            {"name": "Цвет", "value": "черный"},
        ],
    }
    rows = parse_upload("cards.jsonl", json.dumps(record, ensure_ascii=False).encode())
    assert [c.name for c in rows[0].characteristics] == ["Цвет"]
    assert len(rows[0].rejected) == 1


def test_rejected_characteristic_is_reported_in_the_row_result() -> None:
    row = parse_upload("cards.csv", _csv("Батарейка AAA щелочная 1,5 В", f"{LONG_NAME}:Да"))[0]
    result = asyncio.run(process_row(row, _deps()))
    assert any(i.message.startswith("Характеристика не принята") for i in result.issues)
    assert result.status != "ready"


def test_row_waiting_for_an_answer_says_what_to_answer() -> None:
    row = BatchRow(row=1, name="Бумага офисная А4 500 листов")
    result = asyncio.run(process_row(row, _deps()))
    assert result.status == "needs_answer"
    # Белая и цветная бумага — две категории с разными названиями: выбор позиции.
    assert result.note is not None and result.note.startswith("Выберите позицию: ")
    assert "17.12.14.129-00000101" in result.note and "17.12.14.129-00000102" in result.note
    assert "Ни одна не подходит" in result.note


# --- порции пакета: несколько воркеров, повторная доставка, порядок ---


def _job(deps: Deps, n: int, name: str = "Батарейка AAA щелочная 1,5 В") -> str:
    rows = [
        BatchRow(row=i, name=f"{name} {i}", raw={"Наименование": f"{name} {i}"})
        for i in range(1, n + 1)
    ]
    assert deps.batches is not None
    return asyncio.run(deps.batches.create("cards.xlsx", rows))


def test_chunks_cover_every_row_once() -> None:
    chunks = batch_chunks("j", 45, 20)
    assert [(c.start, c.stop) for c in chunks] == [(0, 20), (20, 40), (40, 45)]
    assert batch_chunks("j", 0, 20) == []


def test_chunks_on_two_workers_finish_the_batch_in_original_order() -> None:
    # Два «воркера» над одним Redis, порции вперемешку и одновременно; Redis отдаёт управление
    # внутри каждой команды — гонка read-modify-write была бы видна по счётчику.
    redis = FakeRedis(yield_between=True)
    worker_a, worker_b = _deps(redis), _deps(redis)
    job_id = _job(worker_a, 45)
    chunks = batch_chunks(job_id, 45, 10)

    async def both() -> None:
        await asyncio.gather(
            *(run_chunk(c, worker_a) for c in chunks[::2]),
            *(run_chunk(c, worker_b) for c in reversed(chunks[1::2])),
        )

    asyncio.run(both())
    state = asyncio.run(worker_a.batches.state(job_id))  # type: ignore[union-attr]
    assert state is not None
    assert (state.status, state.total, state.processed) == ("done", 45, 45)
    assert [r.row for r in state.rows] == list(range(1, 46))
    rows = asyncio.run(worker_a.batches.rows(job_id))  # type: ignore[union-attr]
    ws = load_workbook(io.BytesIO(result_workbook(state, rows))).active
    names = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    assert names == [f"Батарейка AAA щелочная 1,5 В {i}" for i in range(1, 46)]


def test_redelivered_chunk_does_not_count_rows_twice() -> None:
    deps = _deps()
    job_id = _job(deps, 12)
    first, second = batch_chunks(job_id, 12, 6)
    assert asyncio.run(run_chunk(first, deps)) == 6
    assert asyncio.run(run_chunk(first, deps)) == 0  # повторная доставка: строки уже есть
    state = asyncio.run(deps.batches.state(job_id))  # type: ignore[union-attr]
    assert state is not None and (state.status, state.processed) == ("running", 6)
    assert sum(state.counts.values()) == 6
    asyncio.run(run_chunk(second, deps))
    state = asyncio.run(deps.batches.state(job_id))  # type: ignore[union-attr]
    assert state is not None and (state.status, state.processed) == ("done", 12)


def test_status_goes_queued_running_done() -> None:
    deps = _deps()
    store = deps.batches
    assert store is not None
    job_id = _job(deps, 3)
    assert asyncio.run(store.state(job_id)).status == "queued"  # type: ignore[union-attr]
    asyncio.run(store.mark_started(job_id))
    assert asyncio.run(store.state(job_id)).status == "running"  # type: ignore[union-attr]
    asyncio.run(run_batch(job_id, deps))
    assert asyncio.run(store.state(job_id)).status == "done"  # type: ignore[union-attr]


def test_empty_file_is_done_at_once() -> None:
    deps = _deps()
    job_id = asyncio.run(deps.batches.create("empty.csv", []))  # type: ignore[union-attr]
    state = asyncio.run(deps.batches.state(job_id))  # type: ignore[union-attr]
    assert state is not None and (state.status, state.total) == ("done", 0)


def test_failed_row_in_a_chunk_does_not_stop_the_others() -> None:
    class BM25BrokenOnRow2:
        def __init__(self, inner: object) -> None:
            self.inner = inner

        def search(self, query: str, *a: object, **kw: object) -> list:
            if " 2" in query:
                raise RuntimeError("индекс сломан")
            return self.inner.search(query, *a, **kw)  # type: ignore[attr-defined]

        def __getattr__(self, name: str) -> object:
            return getattr(self.inner, name)

    deps = _deps()
    deps.bm25 = BM25BrokenOnRow2(deps.bm25)  # type: ignore[assignment]
    job_id = _job(deps, 3)
    asyncio.run(run_batch(job_id, deps))
    state = asyncio.run(deps.batches.state(job_id))  # type: ignore[union-attr]
    assert state is not None and (state.status, state.processed) == ("done", 3)
    assert [r.status == "failed" for r in state.rows] == [False, True, False]


def test_worker_slots_limit_rows_in_flight(monkeypatch: pytest.MonkeyPatch) -> None:
    in_flight = peak = 0

    async def slow_row(row: BatchRow, deps: Deps) -> BatchRowResult:
        nonlocal in_flight, peak
        in_flight += 1
        peak = max(peak, in_flight)
        await asyncio.sleep(0.01)
        in_flight -= 1
        return BatchRowResult(
            row=row.row, name=row.name, status="ready", ktru_code="x", filled=0, issues=[]
        )

    monkeypatch.setattr("app.batch.process_row", slow_row)
    deps = _deps()
    job_id = _job(deps, 6)

    async def go() -> None:
        await run_batch(job_id, deps, asyncio.Semaphore(2))

    asyncio.run(go())
    assert peak == 2  # не по одной (было) и не все шесть сразу


def test_redis_failure_returns_chunk_and_retry_finishes_only_the_rest() -> None:
    class FlakyRedis(FakeRedis):
        fail_after = 2

        async def run_script(self, script: str, keys: list[str], args: list[object]) -> object:
            if "HINCRBY" in script:
                if self.fail_after == 0:
                    raise ConnectionError("Redis недоступен")
                self.fail_after -= 1
            return await super().run_script(script, keys, args)

    redis = FlakyRedis()
    deps = _deps(redis)
    job_id = _job(deps, 4)
    with pytest.raises(ExceptionGroup):
        asyncio.run(run_batch(job_id, deps))  # брокер вернёт порцию в очередь (NACK_ON_ERROR)
    saved = asyncio.run(deps.batches.state(job_id)).processed  # type: ignore[union-attr]
    assert saved == 2
    redis.fail_after = 100
    assert asyncio.run(run_batch(job_id, deps)) is None
    state = asyncio.run(deps.batches.state(job_id))  # type: ignore[union-attr]
    assert state is not None and (state.status, state.processed) == ("done", 4)


def test_result_save_extends_ttl_of_the_whole_batch() -> None:
    redis = FakeRedis()
    deps = _deps(redis)
    job_id = _job(deps, 1)
    redis.ttl.clear()
    asyncio.run(run_batch(job_id, deps))
    assert {k for k, v in redis.ttl.items() if v == 900} == {
        f"batch:{job_id}:meta",
        f"batch:{job_id}:rows",
        f"batch:{job_id}:results",
        f"batch:{job_id}:counts",
        f"batch:{job_id}:index:all",
        f"batch:{job_id}:index:ready",
    }


def _two_rows(store: BatchStore) -> str:
    return asyncio.run(
        store.create("cards.xlsx", [BatchRow(row=1, name="a"), BatchRow(row=2, name="b")])
    )


def _done(row: int) -> BatchRowResult:
    return BatchRowResult(row=row, name="a", status="ready", ktru_code="x", filled=0, issues=[])


def test_batch_waiting_in_the_queue_does_not_expire_before_it_starts() -> None:
    # Пакет стоит в очереди за большим и ещё ничего не записал: срок JOB_TTL_S от загрузки истёк бы
    # раньше, чем до него дойдёт очередь (404 у поставщика, порции потом пропускаются как истёкшие).
    redis = FakeRedis()
    job_id = _two_rows(BatchStore(redis, ttl_s=900, pending_ttl_s=86400))  # type: ignore[arg-type]
    assert redis.ttl[f"batch:{job_id}:meta"] == 86400
    assert redis.ttl[f"batch:{job_id}:rows"] == 86400


def test_batch_ttl_runs_from_the_last_row() -> None:
    redis = FakeRedis()
    store = BatchStore(redis, ttl_s=900, pending_ttl_s=86400)  # type: ignore[arg-type]
    job_id = _two_rows(store)
    keys = [f"batch:{job_id}:{k}" for k in ("meta", "rows", "results", "counts")]
    asyncio.run(store.save_result(job_id, _done(1)))
    assert [redis.ttl[k] for k in keys] == [86400] * 4  # идёт обработка — пакет живёт
    asyncio.run(store.save_result(job_id, _done(2)))
    assert [redis.ttl[k] for k in keys] == [900] * 4  # последняя строка — срок хранения результата


def test_empty_batch_keeps_the_short_ttl() -> None:
    redis = FakeRedis()
    store = BatchStore(redis, ttl_s=900, pending_ttl_s=86400)  # type: ignore[arg-type]
    job_id = asyncio.run(store.create("empty.csv", []))
    assert redis.ttl[f"batch:{job_id}:meta"] == 900


def test_starting_a_chunk_refreshes_the_pending_ttl() -> None:
    redis = FakeRedis()
    store = BatchStore(redis, ttl_s=900, pending_ttl_s=86400)  # type: ignore[arg-type]
    job_id = _two_rows(store)
    asyncio.run(store.save_result(job_id, _done(1)))
    redis.ttl.clear()
    assert asyncio.run(store.mark_started(job_id))
    assert {redis.ttl[f"batch:{job_id}:{k}"] for k in ("meta", "rows", "results", "counts")} == {
        86400
    }


def test_late_delivery_does_not_extend_a_completed_result() -> None:
    redis = FakeRedis()
    store = BatchStore(redis, ttl_s=900)  # type: ignore[arg-type]
    job_id = _two_rows(store)
    asyncio.run(store.save_result(job_id, _done(1)))
    asyncio.run(store.save_result(job_id, _done(2)))
    # Имитируем оставшееся время хранения, повторная доставка его не сбрасывает.
    for key in redis.ttl:
        redis.ttl[key] = 100
    assert not asyncio.run(store.mark_started(job_id))
    assert asyncio.run(store.save_result(job_id, _done(2))) == 2
    assert set(redis.ttl.values()) == {100}


def test_starting_a_gone_batch_does_not_recreate_meta() -> None:
    redis = FakeRedis()
    store = BatchStore(redis, ttl_s=900)  # type: ignore[arg-type]
    job_id = _two_rows(store)
    del redis.hashes[f"batch:{job_id}:meta"]
    assert not asyncio.run(store.mark_started(job_id))
    assert f"batch:{job_id}:meta" not in redis.hashes


def test_result_of_a_gone_batch_is_not_written() -> None:
    # Пакет истёк или удалён: запись строки не создаёт заново results и counts без meta.
    redis = FakeRedis()
    store = BatchStore(redis, ttl_s=900)  # type: ignore[arg-type]
    job_id = _two_rows(store)
    del redis.hashes[f"batch:{job_id}:meta"]
    assert asyncio.run(store.save_result(job_id, _done(1))) == -1
    assert f"batch:{job_id}:results" not in redis.hashes
    assert f"batch:{job_id}:counts" not in redis.hashes


def test_chunk_of_a_gone_batch_does_not_bring_meta_back() -> None:
    redis = FakeRedis()
    deps = _deps(redis)
    job_id = _job(deps, 3)
    del redis.hashes[f"batch:{job_id}:meta"]
    assert asyncio.run(run_batch(job_id, deps)) is None
    assert f"batch:{job_id}:meta" not in redis.hashes
    assert f"batch:{job_id}:results" not in redis.hashes


def test_rows_of_a_chunk_are_kept_when_one_save_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    # Таймаут Redis на записи одной строки не отменяет соседние, которые ещё считаются: их результат
    # записан, порция уходит на повтор только ради упавшей строки.
    first_failed = asyncio.Event()

    class FirstSaveFails(FakeRedis):
        async def run_script(self, script: str, keys: list[str], args: list[object]) -> object:
            if "HINCRBY" in script and not first_failed.is_set():
                first_failed.set()
                raise ConnectionError("Timeout reading from redis:6379")
            return await super().run_script(script, keys, args)

    async def slower_each_row(row: BatchRow, deps: Deps) -> BatchRowResult:
        if row.row > 1:
            await first_failed.wait()  # строка 1 пишется, пока 2–4 ещё считаются
            await asyncio.sleep(0.01)
        return _done(row.row)

    monkeypatch.setattr("app.batch.process_row", slower_each_row)
    redis = FirstSaveFails()
    deps = _deps(redis)
    job_id = _job(deps, 4)
    with pytest.raises(ExceptionGroup):
        asyncio.run(run_batch(job_id, deps))
    assert asyncio.run(deps.batches.state(job_id)).processed == 3  # type: ignore[union-attr]
    assert asyncio.run(run_chunk(batch_chunks(job_id, 4, 4)[0], deps)) == 1
    assert asyncio.run(deps.batches.state(job_id)).processed == 4  # type: ignore[union-attr]


def test_redis_client_waits_longer_than_a_busy_worker_loop() -> None:
    # Строки пакета держат цикл событий воркера (CPU), ответ Redis ждёт своей очереди: короткий
    # таймаут может срабатывать даже при доступном Redis.
    from app.config import Settings
    from app.resources import redis_client

    client = redis_client(Settings(redis_socket_timeout_s=30))
    assert client.connection_pool.connection_kwargs["socket_timeout"] == 30


def test_resources_apply_batch_retention_and_redis_timeout_settings() -> None:
    from app.config import Settings
    from app.resources import Resources, build_deps

    settings = Settings(
        _env_file=None,
        database_url="postgresql+asyncpg://tenderhack:tenderhack@127.0.0.1:5432/tenderhack",
        batch_pending_ttl_s=7200,
        batch_ttl_s=600,
        job_ttl_s=60,
        redis_socket_timeout_s=17,
    )
    resources = Resources.create(settings)
    try:
        deps = build_deps(resources)
        assert deps.batches is not None
        assert deps.batches.pending_ttl_s == 7200
        assert deps.batches.ttl_s == 600
        assert resources.redis.connection_pool.connection_kwargs["socket_timeout"] == 17
    finally:
        asyncio.run(resources.close())


def test_api_state_without_rows_for_polling_big_batches() -> None:
    redis = FakeRedis()
    deps = _deps(redis)
    job_id = _job(deps, 3)
    asyncio.run(run_batch(job_id, deps))
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: deps
    client = TestClient(app)
    light = client.get(f"/api/v1/batch/{job_id}?limit=0").json()
    assert (light["status"], light["processed"], light["rows"]) == ("done", 3, [])
    assert sum(light["counts"].values()) == 3
    assert [r["row"] for r in client.get(f"/api/v1/batch/{job_id}?limit=2").json()["rows"]] == [
        1,
        2,
    ]
    assert len(client.get(f"/api/v1/batch/{job_id}").json()["rows"]) == 3


def test_same_row_saved_twice_counts_once() -> None:
    # Две доставки одной порции успели обработать строку одновременно: счётчики — один раз.
    deps = _deps()
    job_id = _job(deps, 2)
    store = deps.batches
    assert store is not None
    first = BatchRowResult(row=1, name="a", status="ready", ktru_code="x", filled=1, issues=[])
    again = first.model_copy(update={"status": "failed"})
    assert asyncio.run(store.save_result(job_id, first)) == 1
    assert asyncio.run(store.save_result(job_id, again)) == 1
    state = asyncio.run(store.state(job_id))
    assert state is not None
    assert (state.processed, state.counts) == (1, {"ready": 1})
    assert state.rows[0].status == "ready"


def test_upload_as_json_gives_the_same_rows() -> None:
    data = (FAKE / "ste.xlsx").read_bytes()
    as_json = [BatchRow.model_validate_json(r) for r in parse_upload_json("ste.xlsx", data)]
    assert as_json == parse_upload("ste.xlsx", data)


def test_result_file_is_built_in_slices_in_original_order(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.batch._PUSH_SLICE", 4)  # 10 строк — три куска
    deps = _deps(FakeRedis(yield_between=True))
    job_id = _job(deps, 10)
    chunks = batch_chunks(job_id, 10, 3)

    async def go() -> None:
        await asyncio.gather(*(run_chunk(c, deps) for c in reversed(chunks)))

    asyncio.run(go())
    content = asyncio.run(result_file(deps.batches, job_id))  # type: ignore[arg-type]
    ws = load_workbook(io.BytesIO(content)).active
    names = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    assert names == [f"Батарейка AAA щелочная 1,5 В {i}" for i in range(1, 11)]
    assert all(ws.cell(row=i, column=5).value != "не обработано" for i in range(2, 12))


def test_api_row_gives_the_supplier_input_for_the_form() -> None:
    # «Открыть в форме» переносил только наименование: описание и характеристики из файла терялись.
    deps = _deps()
    rows = [
        BatchRow(row=1, name="Бумага офисная А4"),
        BatchRow(
            row=2,
            name="Бумага цветная",
            description="Неон, 500 листов",
            characteristics=[
                CharIn(name="Формат", value="A4"),
                CharIn(name="Плотность", value="80", unit="г/м2"),
            ],
        ),
    ]
    assert deps.batches is not None
    job_id = asyncio.run(deps.batches.create("cards.xlsx", rows))
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: deps
    client = TestClient(app)
    r = client.get(f"/api/v1/batch/{job_id}/rows/2")
    assert r.status_code == 200
    assert r.json() == {
        "row": 2,
        "name": "Бумага цветная",
        "description": "Неон, 500 листов",
        "characteristics": [
            {"name": "Формат", "value": "A4", "unit": None, "origin": None},
            {"name": "Плотность", "value": "80", "unit": "г/м2", "origin": None},
        ],
    }
    assert client.get(f"/api/v1/batch/{job_id}/rows/3").status_code == 404
    assert client.get(f"/api/v1/batch/{job_id}/rows/0").status_code == 404
    assert client.get("/api/v1/batch/nope/rows/1").status_code == 404


def test_empty_file_says_it_is_empty() -> None:
    with pytest.raises(BatchFormatError, match="пустой"):
        parse_upload("cards.csv", b"")


def test_xls_upload_keeps_description_and_characteristics() -> None:
    rows = parse_upload("cards.xls", (FORMATS / "cards.xls").read_bytes())
    assert [r.description for r in rows] == [
        "Плотность 80 г/м2",
        "Количество пар в упаковке 100",
        "",
    ]
    assert [len(r.characteristics) for r in rows] == [3, 2, 2]


def test_process_row_filled_counts_characteristics_not_general_fields() -> None:
    """«Заполнено» в пакете — характеристики: «Производитель» из наименования — поле карточки."""
    from app.catalog import CharacteristicSpec, Position

    maker = CharacteristicSpec(
        name="Производитель", kind="qualitative", values=["Снегурочка"], group="general"
    )
    density = CharacteristicSpec(name="Плотность", kind="quantitative", unit="г/м2")
    cat = Catalog(
        [
            Position(
                code="P",
                name="Бумага для офисной техники",
                okpd2="17.12.14",
                characteristics=[maker, density],
                source="t",
            )
        ],
        "t",
    )
    deps = _deps()
    deps.catalog, deps.bm25 = cat, build_bm25(cat)
    row = BatchRow(row=1, name="Бумага для офисной техники Снегурочка 80 г/м2")
    assert asyncio.run(process_row(row, deps)).filled == 1
