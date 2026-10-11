import asyncio
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.v1 import get_deps
from app.batch import BatchRow, BatchRowResult, run_chunk
from app.jobs import batch_chunks
from app.main import create_app
from tests.fakes import FakeRedis
from tests.test_batch import _deps, _done, _job


def test_cancel_queued_batch_skips_delivery_and_keeps_input(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def forbidden(*args: object) -> None:
        pytest.fail("отменённая строка не должна запускаться")

    monkeypatch.setattr("app.batch.process_row", forbidden)
    redis = FakeRedis()
    deps = _deps(redis)
    store = deps.batches
    assert store is not None
    job_id = _job(deps, 2)
    cancelled = asyncio.run(store.cancel(job_id))
    assert cancelled is not None and cancelled.status == "cancelled"
    assert cancelled.processed == 0
    assert asyncio.run(run_chunk(batch_chunks(job_id, 2, 2)[0], deps)) == 0
    assert len(asyncio.run(store.rows(job_id))) == 2
    assert redis.ttl[f"batch:{job_id}:meta"] == 900
    assert redis.ttl[f"batch:{job_id}:rows"] == 900


def test_cancel_waiting_rows_on_two_workers_keeps_in_flight_results(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def scenario() -> None:
        redis = FakeRedis(yield_between=True)
        workers = [_deps(redis), _deps(redis)]
        store = workers[0].batches
        assert store is not None
        job_id = await store.create("cards.csv", [BatchRow(row=i, name="x") for i in range(1, 5)])
        started: list[int] = []
        both_started, release = asyncio.Event(), asyncio.Event()

        async def processing(row: BatchRow, deps: object) -> BatchRowResult:
            started.append(row.row)
            if len(started) == 2:
                both_started.set()
            await release.wait()
            return _done(row.row)

        monkeypatch.setattr("app.batch.process_row", processing)
        chunks = batch_chunks(job_id, 4, 2)
        tasks = [
            asyncio.create_task(run_chunk(chunk, worker, asyncio.Semaphore(1)))
            for chunk, worker in zip(chunks, workers, strict=True)
        ]
        await asyncio.wait_for(both_started.wait(), 2)
        cancelled = await store.cancel(job_id)
        assert cancelled is not None and cancelled.status == "cancelled"
        release.set()
        assert await asyncio.wait_for(asyncio.gather(*tasks), 2) == [1, 1]
        assert sorted(started) == [1, 3]
        state = await store.state(job_id)
        assert state is not None and (state.status, state.processed) == ("cancelled", 2)
        assert [r.row for r in state.rows] == [1, 3]
        assert state.counts == {"ready": 2}
        assert await run_chunk(chunks[0], workers[0]) == 0
        assert await run_chunk(chunks[1], workers[1]) == 0

    asyncio.run(scenario())


def test_cancel_does_not_affect_another_batch() -> None:
    deps = _deps()
    store = deps.batches
    assert store is not None
    first, other = _job(deps, 1), _job(deps, 1)
    asyncio.run(store.cancel(first))
    assert asyncio.run(run_chunk(batch_chunks(other, 1, 1)[0], deps)) == 1
    assert asyncio.run(store.state(other)).status == "done"  # type: ignore[union-attr]


def test_cancel_again_and_late_results_keep_original_retention_deadline() -> None:
    redis = FakeRedis()
    deps = _deps(redis)
    store = deps.batches
    assert store is not None
    job_id = _job(deps, 2)
    asyncio.run(store.save_result(job_id, _done(1)))
    asyncio.run(store.cancel(job_id))
    for key in redis.ttl:
        redis.ttl[key] = 100
    assert asyncio.run(store.cancel(job_id)).status == "cancelled"  # type: ignore[union-attr]
    assert not asyncio.run(store.mark_started(job_id))
    assert asyncio.run(store.save_result(job_id, _done(2))) == 2
    assert set(redis.ttl.values()) == {100}
    assert asyncio.run(store.state(job_id)).status == "cancelled"  # type: ignore[union-attr]


def test_cancel_completed_batch_is_a_noop() -> None:
    redis = FakeRedis()
    deps = _deps(redis)
    store = deps.batches
    assert store is not None
    job_id = _job(deps, 1)
    asyncio.run(store.save_result(job_id, _done(1)))
    for key in redis.ttl:
        redis.ttl[key] = 100
    result = asyncio.run(store.cancel(job_id))
    assert result is not None and result.status == "done"
    assert set(redis.ttl.values()) == {100}


def test_cancel_missing_batch_does_not_create_keys() -> None:
    redis = FakeRedis()
    store = _deps(redis).batches
    assert store is not None
    assert asyncio.run(store.cancel("missing")) is None
    assert not redis.hashes and not redis.ttl


def test_api_cancel_is_idempotent_and_partial_result_can_be_downloaded() -> None:
    deps = _deps()
    store = deps.batches
    assert store is not None
    job_id = _job(deps, 2)
    asyncio.run(store.save_result(job_id, _done(1)))
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: deps
    client = TestClient(app)
    for _ in range(2):
        response = client.delete(f"/api/v1/batch/{job_id}")
        assert response.status_code == 200
        assert response.json()["status"] == "cancelled"
        assert response.json()["processed"] == 1
    assert client.get(f"/api/v1/batch/{job_id}").json()["status"] == "cancelled"
    download = client.get(f"/api/v1/batch/{job_id}/result.xlsx")
    assert download.status_code == 200
    sheet = load_workbook(io.BytesIO(download.content)).active
    assert sheet is not None
    status_column = next(cell.column for cell in sheet[1] if cell.value == "Статус")
    assert sheet.cell(2, status_column).value == "готово"
    assert sheet.cell(3, status_column).value == "не обработано — пакет прерван"
    assert client.get(f"/api/v1/batch/{job_id}/rows/2").status_code == 200


def test_api_cancel_missing_batch_is_404() -> None:
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: _deps()
    assert TestClient(app).delete("/api/v1/batch/missing").status_code == 404
