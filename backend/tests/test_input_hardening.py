"""Глупый и враждебный ввод поставщика."""

import asyncio
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.api.demo import get_photos, get_submissions
from app.api.v1 import get_deps
from app.batch import (
    BatchFormatError,
    BatchRow,
    BatchRowResult,
    BatchState,
    parse_upload,
    result_workbook,
)
from app.catalog import Catalog
from app.checks.profanity import RULE_PROFANITY
from app.demo import MemorySubmissionRepo
from app.demo.photos import MemoryPhotoRepo
from app.extract import LLMItems
from app.glossary import Glossary
from app.main import create_app
from app.pipeline import AutofillIn, Deps, autofill, check
from app.schemas import MAX_VALUE, CardIn, CharIn
from app.search import build_bm25
from app.session import SessionStore
from app.textclean import clean_text
from tests.fakes import FakeLLM, FakeRedis

CATALOG = Catalog.load([Path(__file__).parent.parent / "app" / "catalog" / "demo" / "demo.json"])
WHITE = "17.12.14.129-00000031"
ZWSP, SHY, NBSP = "\u200b", "\u00ad", "\u00a0"


def _deps() -> Deps:
    return Deps(
        catalog=CATALOG,
        bm25=build_bm25(CATALOG),
        vectors=None,
        llm=FakeLLM([LLMItems(items=[])] * 5),  # type: ignore[arg-type]
        glossary=Glossary([]),
        sessions=SessionStore(FakeRedis(), ttl_s=900),  # type: ignore[arg-type]
        min_score=0.2,
    )


WHITE_CARD = CardIn(
    name="Бумага для офисной техники белая А4",
    ktru_code=WHITE,
    characteristics=[
        CharIn(name="Формат", value="A4"),
        CharIn(name="Цветность", value="Белая"),
        CharIn(name="Масса бумаги площадью 1м2", value=">=80 ... <90", unit="г/м2"),
        CharIn(name="Количество листов в пачке", value="500", unit="шт"),
    ],
)


# --- очистка текста ---


def test_invisible_characters_are_removed_and_spaces_normalized() -> None:
    assert clean_text(f"Бу{ZWSP}мага{NBSP}офи{SHY}сная") == "Бумага офисная"


def test_latin_lookalike_inside_cyrillic_word_becomes_cyrillic() -> None:
    assert clean_text("500 листoв") == "500 листов"


def test_codes_and_signs_stay_as_they_are() -> None:
    assert clean_text("Картридж HP CE285A, A4, № 5") == "Картридж HP CE285A, A4, № 5"


def test_zero_width_spaces_do_not_break_the_search() -> None:
    def codes(name: str) -> list[str]:
        out = asyncio.run(autofill(AutofillIn(name=name), _deps()))
        return [c.code for c in out.ktru.candidates]

    clean = codes("Бумага офисная белая")
    assert clean
    assert codes(f"Бу{ZWSP}мага офи{ZWSP}сная бе{SHY}лая") == clean


# --- лимиты ---


@pytest.mark.parametrize(
    "build",
    [
        lambda: AutofillIn(name="б" * 1001),
        lambda: AutofillIn(name="Бумага", description="б" * 20001),
        lambda: CardIn(name="Бумага", characteristics=[CharIn(name="Ф", value="1")] * 301),
        lambda: CharIn(name="Формат", value="б" * (MAX_VALUE + 1)),
    ],
    ids=["name", "description", "characteristics", "value"],
)
def test_oversized_input_is_rejected(build) -> None:
    with pytest.raises(ValidationError):
        build()


# --- позиция КТРУ ---


def test_card_without_ktru_code_is_blocked() -> None:
    out = check(WHITE_CARD.model_copy(update={"ktru_code": None}), _deps())
    assert out.verdict == "needs_fix"
    assert any(i.field == "ktru_code" and i.severity == "block" for i in out.issues)


def test_unknown_ktru_code_is_blocked() -> None:
    out = check(WHITE_CARD.model_copy(update={"ktru_code": "00.00.00.000-00000000"}), _deps())
    assert any(i.field == "ktru_code" and i.severity == "block" for i in out.issues)


def test_valid_card_is_ready() -> None:
    assert check(WHITE_CARD, _deps()).verdict == "ready"


def test_not_found_autofill_still_reports_issues() -> None:
    out = asyncio.run(autofill(AutofillIn(name="Услуги по уборке, бл*ть, помещений"), _deps()))
    assert out.ktru.final_code is None and out.ktru.question is None
    fields = {(i.field, i.severity) for i in out.issues}
    assert ("ktru_code", "block") in fields
    assert any(i.rule == RULE_PROFANITY for i in out.issues)


# --- стенд заявок ---


def _client() -> TestClient:
    deps = _deps()
    repo = MemorySubmissionRepo()
    app = create_app()
    app.dependency_overrides[get_deps] = lambda: deps
    app.dependency_overrides[get_submissions] = lambda: repo
    photos = MemoryPhotoRepo(repo)
    app.dependency_overrides[get_photos] = lambda: photos
    return TestClient(app)


def test_card_with_blocking_issue_cannot_be_submitted() -> None:
    card = WHITE_CARD.model_copy(update={"name": "Бумага хуёвая А4"})
    r = _client().post("/api/demo/submissions", json={"card": card.model_dump()})
    assert r.status_code == 422
    assert "блокирующ" in r.json()["detail"]


def test_submission_keeps_only_accepted_characteristics_of_the_card() -> None:
    client = _client()
    accepted = [
        {
            "name": "Формат",
            "value": "A4",
            "unit": None,
            "source": None,
            "status": "filled",
            "required": True,
        },
        {
            "name": "ХАК",
            "value": "<script>1</script>",
            "unit": None,
            "source": None,
            "status": "filled",
            "required": False,
        },
    ]
    r = client.post(
        "/api/demo/submissions", json={"card": WHITE_CARD.model_dump(), "accepted": accepted}
    )
    assert r.status_code == 201
    assert [a["name"] for a in r.json()["accepted"]] == ["Формат"]


@pytest.mark.parametrize("note", [None, "", "   "])
def test_return_needs_a_note(note: str | None) -> None:
    client = _client()
    sid = client.post("/api/demo/submissions", json={"card": WHITE_CARD.model_dump()}).json()["id"]
    r = client.post(
        f"/api/demo/submissions/{sid}/decision", json={"decision": "return", "note": note}
    )
    assert r.status_code == 422


# --- массовая загрузка ---


def test_formula_text_stays_text_in_the_result_file() -> None:
    rows = [
        BatchRow(row=2, name="=1+1", raw={"Наименование": '=HYPERLINK("http://example.com","x")'})
    ]
    state = BatchState(
        job_id="j",
        status="done",
        filename="f.xlsx",
        total=1,
        processed=1,
        rows=[
            BatchRowResult(
                row=2, name="=1+1", status="not_found", ktru_code=None, filled=0, issues=[]
            )
        ],
    )
    ws = load_workbook(io.BytesIO(result_workbook(state, rows))).active
    cell = ws.cell(row=2, column=1)
    assert cell.value == '=HYPERLINK("http://example.com","x")'
    assert cell.data_type == "s"
    assert ws.cell(row=2, column=4).data_type == "n"  # «Заполнено» — число


def test_broken_xlsx_gets_a_human_message() -> None:
    with pytest.raises(BatchFormatError, match="повреждён"):
        parse_upload("goods.xlsx", b"not a zip at all")


def test_real_xlsx_is_still_read() -> None:
    data = (Path(__file__).parent / "fixtures" / "fake" / "ste.xlsx").read_bytes()
    assert len(parse_upload("ste.xlsx", data)) == 15


def test_stored_submission_beyond_current_limits_still_loads() -> None:
    # Лимиты — на входе: старая заявка в базе (наименование 10 000 символов)
    # не должна ронять очередь.
    from datetime import UTC, datetime
    from types import SimpleNamespace

    from app.demo.submissions import _row

    old = SimpleNamespace(
        id=1,
        created_at=datetime.now(UTC),
        card={
            "name": "Бумага " + "я" * 10_000,
            "description": "",
            "characteristics": [],
            "ktru_code": None,
        },
        accepted=[],
        status="pending",
        decision_note=None,
    )
    assert len(_row(old).card.name) > 10_000
